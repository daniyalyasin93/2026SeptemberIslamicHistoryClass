# -*- coding: utf-8 -*-
"""
Builds Lecture2.pptx — "The Treasure House: Islam's great historians & books"
A fun primer drawn from the Muqaddima of Tareekh-e-Ummat. Bilingual, 16:9.
Run:  python build_lecture2.py
All death-years are well-established and corroborated by the source text (no invented references).
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK=RGBColor(0x1C,0x23,0x21); TEAL=RGBColor(0x10,0x4A,0x43); TEAL_D=RGBColor(0x0A,0x32,0x2E)
GOLD=RGBColor(0xC4,0x9A,0x45); CREAM=RGBColor(0xFA,0xF7,0xF0); CREAM_D=RGBColor(0xEF,0xE9,0xDB)
GREY=RGBColor(0x8C,0x8C,0x86); WHITE=RGBColor(0xFF,0xFF,0xFF); MAROON=RGBColor(0x7A,0x2E,0x2E)
AR="Traditional Arabic"; EN="Georgia"; SANS="Segoe UI"

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
SW,SH=prs.slide_width,prs.slide_height; BLANK=prs.slide_layouts[6]

def slide(): return prs.slides.add_slide(BLANK)
def bg(s,c=CREAM):
    s.background.fill.solid(); s.background.fill.fore_color.rgb=c
def rect(s,x,y,w,h,fill=None,line=None,lw=1.0,shape=MSO_SHAPE.RECTANGLE):
    sp=s.shapes.add_shape(shape,x,y,w,h)
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(lw)
    sp.shadow.inherit=False; return sp
def set_cs(run,name):
    rPr=run._r.get_or_add_rPr(); cs=rPr.find(qn('a:cs'))
    if cs is None: cs=rPr.makeelement(qn('a:cs'),{}); rPr.append(cs)
    cs.set('typeface',name)
def rtl(p): p._p.get_or_add_pPr().set('rtl','1')
def tb(s,x,y,w,h,anchor=MSO_ANCHOR.TOP):
    box=s.shapes.add_textbox(x,y,w,h); tf=box.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=Pt(4); tf.margin_top=tf.margin_bottom=Pt(2)
    return box,tf
def para(tf,text,size=20,color=INK,bold=False,font=SANS,align=PP_ALIGN.LEFT,first=False,
         space_after=6,italic=False,arabic=False):
    p=tf.paragraphs[0] if first and tf.paragraphs[0].text=="" else tf.add_paragraph()
    p.alignment=align; p.space_after=Pt(space_after)
    r=p.add_run(); r.text=text; f=r.font
    f.size=Pt(size); f.bold=bold; f.italic=italic; f.color.rgb=color; f.name=font
    if arabic: set_cs(r,AR); rtl(p)
    return p,r
def kicker(s,text,color=GOLD):
    _,tf=tb(s,Inches(0.7),Inches(0.45),Inches(11.9),Inches(0.5))
    para(tf,text.upper(),size=14,color=color,bold=True,font=SANS,first=True,space_after=0)
def goldbar(s,x=0.7,y=0.86,w=1.6): rect(s,Inches(x),Inches(y),Inches(w),Pt(3),fill=GOLD)
def headline(s,text,y=0.95,size=32,color=TEAL_D):
    _,tf=tb(s,Inches(0.7),Inches(y),Inches(11.9),Inches(1.2))
    para(tf,text,size=size,color=color,bold=True,font=EN,first=True,space_after=0)
def notes(s,t): s.notes_slide.notes_text_frame.text=t
def footer(s,n):
    _,tf=tb(s,Inches(11.2),Inches(7.0),Inches(1.9),Inches(0.4))
    para(tf,f"Lesson 2  ·  {n}",size=10,color=GREY,align=PP_ALIGN.RIGHT,first=True,space_after=0)

def card_row(s, items, y, h=2.35, accent=TEAL, title_size=19, desc_size=15):
    """items: list of (name, date_or_sub, desc)"""
    n=len(items); margin=0.7; gap=0.3
    total=13.333-2*margin; w=(total-gap*(n-1))/n; x=margin
    for (name,sub,desc) in items:
        rect(s,Inches(x),Inches(y),Inches(w),Inches(h),fill=CREAM_D,line=accent,lw=1.0)
        rect(s,Inches(x),Inches(y),Inches(w),Pt(5),fill=GOLD)
        _,tf=tb(s,Inches(x+0.18),Inches(y+0.16),Inches(w-0.36),Inches(h-0.3))
        para(tf,name,size=title_size,bold=True,color=TEAL_D,first=True,space_after=2)
        if sub: para(tf,sub,size=13,color=GOLD,bold=True,space_after=6)
        para(tf,desc,size=desc_size,color=INK,space_after=0)
        x+=w+gap

def section(title,sub_ar):
    s=slide(); bg(s,TEAL_D)
    rect(s,Inches(0.7),Inches(3.0),Inches(2.0),Pt(4),fill=GOLD)
    _,tf=tb(s,Inches(0.7),Inches(3.2),Inches(12),Inches(2.2))
    para(tf,title,size=42,color=WHITE,bold=True,font=EN,first=True,space_after=8)
    para(tf,sub_ar,size=26,color=GOLD,font=AR,arabic=True,space_after=0)
    return s

# ===== 1 TITLE
s=slide(); bg(s,TEAL_D)
rect(s,0,Inches(3.05),SW,Inches(1.4),fill=TEAL); rect(s,Inches(0.7),Inches(2.78),Inches(2.2),Pt(4),fill=GOLD)
_,tf=tb(s,Inches(0.7),Inches(1.5),Inches(12),Inches(1.2))
para(tf,"LESSONS FROM HISTORY  ·  LESSON 2",size=18,color=GOLD,bold=True,font=SANS,first=True,space_after=2)
para(tf,"تاریخ سے سبق",size=20,color=CREAM,font=AR,arabic=True,space_after=0)
_,tf=tb(s,Inches(0.7),Inches(3.15),Inches(12),Inches(1.25),anchor=MSO_ANCHOR.MIDDLE)
para(tf,"The Treasure House",size=44,color=WHITE,bold=True,font=EN,first=True,space_after=0)
_,tf=tb(s,Inches(0.7),Inches(4.7),Inches(12),Inches(1))
para(tf,"اسلام کے عظیم مؤرخین اور تاریخ کی عظیم کتب",size=26,color=CREAM,font=AR,arabic=True,first=True,space_after=4)
para(tf,"Who saved our history — and how we know we can trust it",size=16,color=GOLD,font=SANS,space_after=0)
notes(s,"Recap last week in 30 sec: 'Pichle hafte humne dekha tareekh ke USOOL — sand aur asma al-rijal. "
        "Aaj hum un LOGON se milenge jinhone yeh sab mehfooz kiya — aur un KITABON se jin par hum bharosa karte hain.' "
        "Keep it celebratory, like meeting heroes.")
footer(s,"Title")

# ===== 2 HOOK
s=slide(); bg(s,GOLD)
_,tf=tb(s,Inches(1.1),Inches(2.2),Inches(11.1),Inches(3.1),anchor=MSO_ANCHOR.MIDDLE)
para(tf,"It is 1400 years later.",size=30,color=TEAL_D,bold=True,font=EN,first=True,space_after=14,align=PP_ALIGN.CENTER)
para(tf,"How do we know ANY of this actually happened?",size=34,color=INK,bold=True,font=EN,align=PP_ALIGN.CENTER,space_after=14)
para(tf,"Because an army of scholars spent their lives saving it — with rules no other nation had.",
     size=22,color=TEAL_D,font=EN,align=PP_ALIGN.CENTER,italic=True,space_after=0)
notes(s,"Pose the question to the room and pause. 'Socho — Rome aur Yunan ki tareekh ka silsila ustad se "
        "mehroom hai. Hamare paas sand hai. Kyun? Kyunki kuch log uth khade hue.' Then go to the two pillars.")
footer(s,"Hook")

# ===== 3 TWO PILLARS
s=slide(); bg(s); kicker(s,"Islamic history stands on two pillars — both born from the science of hadith")
goldbar(s); headline(s,"Two pillars")
card_row(s,[
 ("Seerah-nigārī","سیرت و مغازی","Recording the Prophet's ﷺ life in time-order. Grew out of hadith into its own science — the root of all Islamic history-writing."),
 ("Fann al-Rijāl","علم اسماء الرجال","The biographies of the narrators. Built to protect hadith AND history — so any report can be graded and any liar caught."),
], y=2.5, h=3.0, title_size=24, desc_size=19)
notes(s,"Keep simple: 'Do buniyaadein: ek, seerat-nigari — Nabi ﷺ ki zindagi tarteeb se. Doosri, fann-e-rijaal — "
        "har raavi ka koi firishta nahi, uska poora hisaab.' Ref: Introduction.pdf p.43-44.")
footer(s,"Pillars")

# ===== 4 FOUR STAGES (house)
s=slide(); bg(s); kicker(s,"Every science is built like a house — in four stages")
goldbar(s); headline(s,"Building the house of history",size=30)
stages=[("1 · TĀSĪS","تاسیس","Lay the foundation — mark out the new science",GREY),
        ("2 · TADWĪN","تدوین","Raise the walls — gather the raw material",TEAL),
        ("3 · TANQĪH","تنقیح و تہذیب","Plaster & clean — weed out the weak, order the sound",TEAL_D),
        ("4 · TAKMĪL","تکمیل","Paint & furnish — perfect, simplify, publish",GOLD)]
margin=0.7; gap=0.25; n=4; total=13.333-2*margin; w=(total-gap*(n-1))/n; x=margin
for (t,ar,d,col) in stages:
    rect(s,Inches(x),Inches(2.4),Inches(w),Inches(3.2),fill=CREAM_D,line=col,lw=1.5)
    rect(s,Inches(x),Inches(2.4),Inches(w),Inches(0.55),fill=col)
    _,tf=tb(s,Inches(x+0.15),Inches(2.46),Inches(w-0.3),Inches(0.5))
    para(tf,t,size=16,bold=True,color=WHITE,first=True,space_after=0)
    _,tf=tb(s,Inches(x+0.15),Inches(3.1),Inches(w-0.3),Inches(2.4))
    para(tf,ar,size=22,bold=True,color=col,font=AR,arabic=True,first=True,space_after=8)
    para(tf,d,size=15,color=INK,space_after=0)
    x+=w+gap
_,tf=tb(s,Inches(0.7),Inches(5.9),Inches(11.9),Inches(1.0))
para(tf,"Islamic history: founded before Islam · collected in the 2nd–4th centuries AH · refined in the 7th · perfected in the 8th (Ibn Khaldūn).",
     size=18,italic=True,color=TEAL,first=True,space_after=0)
notes(s,"The book uses this exact 'building a house' metaphor — lean into it, it's memorable. "
        "Ref: Introduction.pdf p.41-42, 45.")
footer(s,"4 stages")

# ===== 5 SECTION giants
section("Meet the giants","عظیم مؤرخین سے ملیے")
notes(prs.slides[-1],"Treat the next slides like a hall of fame / trading cards. Energy up. "
                     "You don't need every date memorised — the cards carry them.")
footer(prs.slides[-1],"Section")

# ===== 6 Seerah pioneers
s=slide(); bg(s); kicker(s,"The pioneers — who first wrote the Prophet's ﷺ life down")
goldbar(s); headline(s,"The seerah pioneers",size=30)
card_row(s,[
 ("Ibn Shihāb al-Zuhrī","d. 124 AH","Compiled the first book of maghāzī — under a caliph's order."),
 ("Ibn Isḥāq","d. 151 AH","The great seerah collector; gathered everything (not always with strict chains)."),
 ("Ibn Hishām","d. 213 AH","Refined Ibn Isḥāq into Sīrat Ibn Hishām — the most popular seerah source of all."),
], y=2.6, h=3.0)
notes(s,"Story-thread: 'Pehle yeh sab seenon mein tha. Phir Zuhri ne likha, Ibn Ishaq ne jama kiya, "
        "Ibn Hisham ne saaf kar ke woh kitab di jo aaj tak sab se mashhoor hai.' Ref: Introduction.pdf p.43-44.")
footer(s,"Giants 1")

# ===== 7 Waqidi & Ibn Sa'd
s=slide(); bg(s); kicker(s,"The vivid one and the careful one — teacher & student")
goldbar(s); headline(s,"Wāqidī & Ibn Saʿd",size=30)
card_row(s,[
 ("Al-Wāqidī","d. 207 AH","A judge of Baghdad and a giant memory. His Maghāzī is vivid and detailed — but he gathered weak material too, so scholars handle him with care."),
 ("Ibn Saʿd","d. 230 AH","Wāqidī's student — but stricter. His Ṭabaqāt al-Kubrā is a huge, ordered treasury of the Companions' lives that no historian can do without."),
], y=2.6, h=3.2, title_size=23, desc_size=18)
notes(s,"Nice contrast to make it human: 'Ustad rangeen tha lekin thoda be-ehtiyaat; shagird zyada muhtaat nikla.' "
        "Drives home WHY we grade narrators. Ref: Introduction.pdf p.45.")
footer(s,"Giants 2")

# ===== 8 encyclopedists + caution
s=slide(); bg(s); kicker(s,"The first true tārīkhs — and a caution")
goldbar(s); headline(s,"The encyclopedists",size=30)
card_row(s,[
 ("Khalīfa b. Khayyāṭ","d. 240 AH","His annalistic Tārīkh is arguably the first proper Muslim history — sober, year-by-year, reliable narrators."),
 ("Al-Balādhurī","d. 279 AH","Futūḥ al-Buldān & Ansāb al-Ashrāf — careful, city-by-city conquests; a primary source."),
], y=2.5, h=2.4, title_size=21, desc_size=17)
rect(s,Inches(0.7),Inches(5.2),Inches(11.93),Inches(1.5),fill=CREAM_D,line=MAROON,lw=1.5)
_,tf=tb(s,Inches(1.0),Inches(5.35),Inches(11.3),Inches(1.2))
para(tf,"⚠  Know your author's lens:",size=16,bold=True,color=MAROON,first=True,space_after=4)
para(tf,"Some early historians — e.g. al-Yaʿqūbī (3rd c.) and al-Masʿūdī (d. 346 AH) — were Shīʿa; parts of their work "
        "tilt against the Companions. We read them aware, not blind.",size=17,color=INK,space_after=0)
notes(s,"This 'know the lens' point is gold for the controversy brief — it teaches the audience to ask "
        "WHO wrote a thing. Ref: Introduction.pdf p.46, 68.")
footer(s,"Giants 3")

# ===== 9 SECTION mother books
section("The five mother-books","تاریخ کی پانچ اہم کتب (موسوعات)")
footer(prs.slides[-1],"Section")

# ===== 10 mother-books overview shelf
s=slide(); bg(s); kicker(s,"The five great encyclopedic histories everything else draws on")
goldbar(s); headline(s,"The library shelf",size=30)
card_row(s,[
 ("Tārīkh al-Ṭabarī","Ṭabarī · d. 310","The great well: every report with its chain."),
 ("Al-Kāmil","Ibn al-Athīr · d. 630","Year-by-year, distilled and readable."),
 ("Tārīkh al-Islām","Dhahabī · d. 748","The most careful sieve of all."),
 ("Al-Bidāya wa'l-Nihāya","Ibn Kathīr · d. 774","Reason + chains; the reader-friendly one."),
 ("Tārīkh Ibn Khaldūn","Ibn Khaldūn · d. 808","Its Muqaddima founded social science."),
], y=2.7, h=3.0, title_size=16, desc_size=14)
notes(s,"Present like a bookshelf reveal. 'Yeh paanch maa-kitaben hain — baaqi sab tareekhein inhi se peeti hain.' "
        "Ref: Introduction.pdf p.69-78.")
footer(s,"Mother-books")

# ===== 11 Tabari deep + fun fact
s=slide(); bg(s); kicker(s,"The single most important source — and a famous mix-up")
goldbar(s); headline(s,"Tārīkh al-Ṭabarī",size=30)
_,tf=tb(s,Inches(0.7),Inches(2.05),Inches(11.9),Inches(2.7))
for t in [
 "Abū Jaʿfar Muḥammad b. Jarīr al-Ṭabarī (d. 310 AH) — himself a master mufassir, muḥaddith and faqīh.",
 "He wrote tārīkh bi'l-riwāyah: he records each report WITH its chain, changing nothing — so reading him is like reading the lost early sources verbatim.",
 "That is also its catch: he gathered sound AND weak reports, leaving the sifting to us. Don't read it raw without training.",
]:
    para(tf,"•  "+t,size=20,first=(t.startswith("Abū")),space_after=12)
rect(s,Inches(0.7),Inches(5.0),Inches(11.93),Inches(1.7),fill=CREAM_D,line=GOLD,lw=1.5)
_,tf=tb(s,Inches(1.0),Inches(5.15),Inches(11.3),Inches(1.4))
para(tf,"🎯  FUN FACT — two men, one name:",size=16,bold=True,color=TEAL,first=True,space_after=4)
para(tf,"There were TWO Abū Jaʿfar Muḥammad b. Jarīr al-Ṭabarīs. Ours (grandson of Yazīd) is the Sunnī imam; "
        "the other (descendant of Rustam) was Shīʿa. The name-clash has fooled many — including into calling our Ṭabarī a Shīʿa.",
     size=16,color=INK,space_after=0)
notes(s,"The namesake fun-fact is true and disarming — and pre-empts a common objection (that Tabari was Shia). "
        "Hafiz al-Dhahabi explicitly defended him. Ref: Introduction.pdf p.69-70.")
footer(s,"Ṭabarī")

# ===== 12 the other four strengths
s=slide(); bg(s); kicker(s,"…and how the other four improved on it")
goldbar(s); headline(s,"From raw well to refined history",size=30)
_,tf=tb(s,Inches(0.7),Inches(2.1),Inches(11.9),Inches(4.5))
for n,t in [
 ("Al-Kāmil — Ibn al-Athīr (d. 630):","took the gist, dropped the clutter, ran it year-by-year across the whole world."),
 ("Tārīkh al-Islām — al-Dhahabī (d. 748):","the most careful selection; merges history + narrators + generations in one."),
 ("Al-Bidāya wa'l-Nihāya — Ibn Kathīr (d. 774):","weighs chains AND reason; careful with the seerah; the friendliest read."),
 ("Tārīkh Ibn Khaldūn (d. 808):","its Muqaddima theorised why nations rise and fall — the birth of sociology."),
]:
    p,_=para(tf,n+" ",size=20,bold=True,color=TEAL,first=(n.startswith("Al-Kāmil")),space_after=12)
    r=p.add_run(); r.text=t; r.font.size=Pt(20); r.font.color.rgb=INK; r.font.name=SANS
notes(s,"Show the ARC: raw (Tabari) → distilled (Ibn Athir) → sifted (Dhahabi) → reasoned (Ibn Kathir) → "
        "theorised (Ibn Khaldun). That arc IS the 4 stages in action. Ref: Introduction.pdf p.73-78.")
footer(s,"The four")

# ===== 13 hidden gems
s=slide(); bg(s); kicker(s,"Two treasures most people overlook")
goldbar(s); headline(s,"The hidden gems",size=30)
card_row(s,[
 ("Al-Muntaẓam","Ibn al-Jawzī · d. 597 AH","By a standard of rigour, finer than Ṭabarī in places — yet far less famous."),
 ("Mir'āt al-Zamān","Sibṭ Ibn al-Jawzī · d. 654 AH","His grandson's work; in craft, it outdoes al-Kāmil. Long lost, recently reassembled."),
], y=2.6, h=2.7, title_size=22, desc_size=18)
_,tf=tb(s,Inches(0.7),Inches(5.6),Inches(11.9),Inches(0.9))
para(tf,"Dhahabī and Ibn Kathīr quietly leaned on both. Worth knowing the gems, not just the famous five.",
     size=18,italic=True,color=TEAL,first=True,space_after=0)
notes(s,"Optional/bonus — drop if short on time. Ref: Introduction.pdf p.78.")
footer(s,"Gems")

# ===== 14 SECTION detective
section("The detective science","جعلی روایات کیسے پکڑی جاتی ہیں")
footer(prs.slides[-1],"Section")

# ===== 15 Shafi'i wine story
s=slide(); bg(s); kicker(s,"How the rules catch a forged story")
goldbar(s); headline(s,"The case of the “wine cup”",size=30)
_,tf=tb(s,Inches(0.7),Inches(2.1),Inches(11.9),Inches(4.6))
para(tf,"A report circulated: Maʾmūn al-Rashīd once tested Imām al-Shāfiʿī with so much wine it would floor any man — yet it had no effect on him.",
     size=22,first=True,space_after=14)
para(tf,"Ḥāfiẓ Ibn Ḥajar dismantled it in Lisān al-Mīzān, on two grounds:",size=22,bold=True,color=TEAL,space_after=10)
para(tf,"•  History: al-Shāfiʿī meeting Maʾmūn is not even established — the chain fails.",size=21,space_after=8)
para(tf,"•  Reason (dirāyah): al-Shāfiʿī said he'd avoid even COLD WATER if it dulled his mind — so a wine-cup tale is absurd.",size=21,space_after=0)
notes(s,"A fresh, fun example (different from last week's forged document). Punch the ending: 'Jo aadmi thanda paani "
        "is liye chhor de ke kahin zehn par asar na ho — uske baare mein sharab ki kahani?!' "
        "Ref: Introduction.pdf p.57; Lisān al-Mīzān (Ibn Ḥajar).")
footer(s,"Detective")

# ===== 16 why it matters
s=slide(); bg(s,TEAL)
rect(s,Inches(0.7),Inches(1.5),Pt(5),Inches(4.4),fill=GOLD)
_,tf=tb(s,Inches(1.3),Inches(1.7),Inches(11.2),Inches(4.2),anchor=MSO_ANCHOR.MIDDLE)
para(tf,"This is why we can trust what we'll say.",size=32,color=WHITE,bold=True,font=EN,first=True,space_after=18)
for t in ["These books and these rules are the ground under every claim in this course.",
          "When we reach the hard chapters, we'll stand on this — not on guesswork, not on novels, not on films.",
          "No other civilisation built a machine like this to guard its own past."]:
    para(tf,"•  "+t,size=22,color=CREAM,space_after=12)
notes(s,"Tie firmly back to Lesson 1's contrast. This is the payoff: rigour = trust. "
        "Then the take-home.")
footer(s,"Why")

# ===== 17 close
s=slide(); bg(s,TEAL_D)
rect(s,Inches(0.7),Inches(1.2),Inches(2.0),Pt(4),fill=GOLD)
_,tf=tb(s,Inches(0.7),Inches(1.4),Inches(11.9),Inches(2.2))
para(tf,"Next week",size=18,color=GOLD,bold=True,first=True,space_after=8)
para(tf,"The story begins: Abu Bakr ؓ",size=38,color=WHITE,bold=True,font=EN,space_after=6)
para(tf,"خلافتِ راشدہ اور سیدنا ابوبکر صدیق ؓ",size=26,color=CREAM,font=AR,arabic=True,space_after=0)
_,tf=tb(s,Inches(0.7),Inches(4.4),Inches(11.9),Inches(2.4))
para(tf,"We place our first bright peg on the timeline — armed now with the books and the rules to read it safely.",
     size=22,color=CREAM,first=True,space_after=18)
para(tf,"Take-home: name ONE mother-book and ONE thing that makes Islamic history different from the rest.",
     size=22,color=GOLD,bold=True,italic=True,space_after=0)
notes(s,"End on the take-home (answers: a mother-book = Tabari/al-Kamil/etc.; the difference = isnad/asma al-rijal). "
        "Keep the timeline banner visible; add the 'books' as a small label on it if you like.")
footer(s,"Close")

# ===== 18 sources
s=slide(); bg(s); kicker(s,"Where this comes from"); goldbar(s); headline(s,"Sources",size=28)
_,tf=tb(s,Inches(0.7),Inches(2.1),Inches(11.9),Inches(4.6))
for t in [
 "Primary course text — Tareekh-e-Ummat, Maulana Muhammad Ismail Rehan, Vol. 1 (Muqaddima).",
 "Two foundations & the development of historiography — muqaddima pp. 43–48.",
 "The four stages — muqaddima pp. 41–42, 45.",
 "The five mother-books & the hidden gems — muqaddima pp. 69–78.",
 "Ṭabarī namesake & defence — muqaddima pp. 69–70 (Mīzān al-Iʿtidāl, al-Dhahabī).",
 "The forged 'wine' report — muqaddima p. 57 (Lisān al-Mīzān, Ibn Ḥajar).",
]:
    para(tf,"•  "+t,size=18,first=(t.startswith("Primary")),space_after=11)
notes(s,"Keep for credibility / handouts. All dates above are well-established and match the source text.")
footer(s,"Sources")

import os
out=os.path.join(os.path.dirname(os.path.abspath(__file__)),"Lecture2.pptx")
prs.save(out)
print("Saved:",out,"·",len(prs.slides._sldIdLst),"slides")
