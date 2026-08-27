"""Shared python-pptx kit for the lecture decks — English-first, picture-led.

    import sys, os
    sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "series"))
    from deckkit import *

DESIGN RULES, settled with Daniyal 2026-08-27:

1. **English carries the slide.** Headlines, structure and labels in English. Urdu appears only
   where it is doing real work, and **always with an English rendering beside it**. Arabic
   quotations stay in Naskh with an English translation under them.
2. **Latin digits on slides, always.** Urdu-Indic digits mixed with Latin punctuation get
   reordered by the bidi algorithm and render as nonsense. This was a real bug in the first decks.
3. **No bare-text slides.** Every content slide carries a map, the timeline strip, a large
   number, or a quotation. If it is only a line of text, it is not a slide.
4. **Urdu font is `Noto Nastaliq Urdu`** — the previously used Jameel cut installed on this
   machine is the *Kasheeda* variant, which elongates letters and renders badly.
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

TEAL = RGBColor(0x10, 0x4A, 0x43)
DARK = RGBColor(0x0A, 0x32, 0x2E)
GOLD = RGBColor(0xC4, 0x9A, 0x45)
CREAM = RGBColor(0xFA, 0xF7, 0xF0)
INK = RGBColor(0x1C, 0x23, 0x21)
MUTED = RGBColor(0x5A, 0x68, 0x62)
MAROON = RGBColor(0x7A, 0x2E, 0x2E)

UR = "Noto Nastaliq Urdu"
AR = "Traditional Arabic"
EN = "Georgia"
SANS = "Segoe UI"

W, H = Inches(13.333), Inches(7.5)
VIS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "visuals")


def set_cs(run, face):
    rPr = run._r.get_or_add_rPr()
    cs = rPr.find(qn("a:cs"))
    if cs is None:
        cs = rPr.makeelement(qn("a:cs"), {})
        rPr.append(cs)
    cs.set("typeface", face)


def deck():
    p = Presentation()
    p.slide_width, p.slide_height = W, H
    return p


def blank(prs, bg=CREAM):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    r = s.shapes.add_shape(1, 0, 0, W, H)
    r.fill.solid(); r.fill.fore_color.rgb = bg; r.line.fill.background()
    r.shadow.inherit = False
    return s


def text(slide, txt, x, y, w, h, size=28, color=INK, font=EN, cs=None,
         bold=False, align=PP_ALIGN.LEFT, is_rtl=False, spacing=1.0, italic=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(txt.split("\n")):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.line_spacing = spacing
        run = par.add_run(); run.text = line
        f = run.font
        f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
        f.color.rgb = color
        if cs:
            set_cs(run, cs)
        par.alignment = align
        if is_rtl:
            par._p.get_or_add_pPr().set("rtl", "1")
            par.alignment = PP_ALIGN.RIGHT
    return tb


def rule(slide, x, y, w, color=GOLD, thick=Pt(2.5)):
    ln = slide.shapes.add_shape(1, x, y, w, Emu(int(thick)))
    ln.fill.solid(); ln.fill.fore_color.rgb = color; ln.line.fill.background()
    ln.shadow.inherit = False


def band(slide, y, h, color=DARK, x=0, w=W):
    r = slide.shapes.add_shape(1, x, y, w, h)
    r.fill.solid(); r.fill.fore_color.rgb = color; r.line.fill.background()
    r.shadow.inherit = False
    return r


def note(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def eyebrow(slide, label, color=TEAL, y=Inches(0.46)):
    text(slide, label.upper(), Inches(0.85), y, Inches(9), Inches(0.4),
         size=12, color=color, font=SANS, bold=True)


def strip(slide, which, y=Inches(6.62)):
    """The timeline strip — where tonight sits on the whole line. On every content slide."""
    p = os.path.join(VIS, "%s.png" % which)
    if os.path.exists(p):
        slide.shapes.add_picture(p, Inches(0.6), y, width=Inches(12.13))


# ------------------------------------------------------------------ slide patterns
def title_slide(prs, en, ur, sub, meta):
    s = blank(prs, DARK)
    text(s, en, Inches(0.9), Inches(1.75), Inches(11.5), Inches(1.5),
         size=62, color=CREAM, font=EN, bold=False)
    if ur:
        text(s, ur, Inches(0.9), Inches(3.15), Inches(11.5), Inches(1.0),
             size=40, color=GOLD, font=UR, cs=UR, align=PP_ALIGN.RIGHT, is_rtl=True, spacing=1.9)
    rule(s, Inches(0.9), Inches(4.55), Inches(2.6))
    text(s, sub, Inches(0.9), Inches(4.85), Inches(11.5), Inches(0.8),
         size=24, color=CREAM, font=EN, italic=True)
    text(s, meta, Inches(0.9), Inches(6.3), Inches(11.5), Inches(0.5),
         size=13, color=GOLD, font=SANS, bold=True)
    return s


def map_slide(prs, image, headline, sub=None, line=None, caption=None):
    """Full-bleed map with a headline band. The map IS the slide."""
    s = blank(prs, CREAM)
    p = os.path.join(VIS, "%s.png" % image)
    if os.path.exists(p):
        s.shapes.add_picture(p, Inches(0), Inches(0.95), width=W)
    band(s, 0, Inches(0.95), DARK)
    text(s, headline, Inches(0.7), Inches(0.16), Inches(8.6), Inches(0.7),
         size=30, color=CREAM, font=EN)
    if sub:
        text(s, sub, Inches(9.4), Inches(0.28), Inches(3.3), Inches(0.5),
             size=15, color=GOLD, font=SANS, align=PP_ALIGN.RIGHT)
    if line:
        strip(s, line)
    return s


def quote_slide(prs, arabic, english, cite, label=None, urdu=None):
    """A sourced quotation. Arabic in Naskh, ENGLISH translation beneath, reference under that.

    The attribution carries the claim — the speaker never does.
    """
    s = blank(prs, CREAM)
    if label:
        eyebrow(s, label)
    y = 1.45
    text(s, arabic, Inches(0.7), Inches(y), Inches(11.9), Inches(1.7),
         size=36, color=DARK, font=AR, cs=AR, align=PP_ALIGN.CENTER, is_rtl=True, spacing=1.4)
    y += 1.85
    text(s, english, Inches(1.3), Inches(y), Inches(10.7), Inches(1.5),
         size=25, color=INK, font=EN, align=PP_ALIGN.CENTER, spacing=1.35, italic=True)
    y += 1.5
    if urdu:
        text(s, urdu, Inches(1.0), Inches(y), Inches(11.3), Inches(1.0),
             size=24, color=MUTED, font=UR, cs=UR, align=PP_ALIGN.CENTER, is_rtl=True, spacing=1.9)
        y += 1.05
    rule(s, Inches(5.9), Inches(y + 0.12), Inches(1.5), GOLD, Pt(1.5))
    text(s, cite, Inches(0.7), Inches(y + 0.3), Inches(11.9), Inches(0.5),
         size=14, color=TEAL, font=SANS, align=PP_ALIGN.CENTER)
    return s


def statement_slide(prs, big, sub=None, line=None, kicker=None, bg=CREAM, fg=None):
    """One large statement. Used sparingly — for a turn in the story."""
    s = blank(prs, bg)
    fg = fg or (CREAM if bg in (DARK, TEAL) else TEAL)
    if kicker:
        eyebrow(s, kicker, GOLD if bg in (DARK, TEAL) else TEAL)
    tb = text(s, big, Inches(0.9), Inches(2.15), Inches(11.5), Inches(2.2),
              size=52, color=fg, font=EN, spacing=1.2)
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    if sub:
        text(s, sub, Inches(0.9), Inches(4.55), Inches(11.5), Inches(1.0),
             size=23, color=(GOLD if bg in (DARK, TEAL) else MUTED), font=EN, italic=True, spacing=1.3)
    if line:
        strip(s, line)
    return s


def person_slide(prs, name_en, name_ur, laqab_en, laqab_ur, dates, did, incident, cite, line=None):
    """A تعارف panel. English carries it; the Urdu name sits beside, not instead."""
    s = blank(prs, CREAM)
    eyebrow(s, "who you are meeting")
    text(s, name_en, Inches(0.85), Inches(1.05), Inches(8.4), Inches(0.9),
         size=40, color=TEAL, font=EN)
    text(s, name_ur, Inches(8.6), Inches(1.0), Inches(3.9), Inches(0.9),
         size=30, color=TEAL, font=UR, cs=UR, align=PP_ALIGN.RIGHT, is_rtl=True, spacing=1.9)
    if laqab_en:
        band(s, Inches(1.98), Inches(0.62), GOLD, Inches(0.85), Inches(11.6))
        text(s, laqab_en, Inches(1.05), Inches(2.06), Inches(7.0), Inches(0.5),
             size=19, color=DARK, font=SANS, bold=True)
        text(s, laqab_ur, Inches(8.0), Inches(2.0), Inches(4.3), Inches(0.55),
             size=20, color=DARK, font=UR, cs=UR, align=PP_ALIGN.RIGHT, is_rtl=True, spacing=1.7)
    if dates:
        text(s, dates, Inches(0.85), Inches(2.88), Inches(11.6), Inches(0.45),
             size=15, color=MUTED, font=SANS, bold=True)
    # generous box — two lines at this size overflowed into the divider below
    text(s, did, Inches(0.85), Inches(3.4), Inches(11.6), Inches(1.3),
         size=22, color=INK, font=EN, spacing=1.3)
    rule(s, Inches(0.85), Inches(4.72), Inches(1.6), GOLD, Pt(2))
    text(s, "ONE INCIDENT", Inches(0.85), Inches(4.9), Inches(4), Inches(0.35),
         size=11, color=MUTED, font=SANS, bold=True)
    text(s, incident, Inches(0.85), Inches(5.26), Inches(11.6), Inches(1.25),
         size=20, color=DARK, font=EN, spacing=1.3, italic=True)
    text(s, cite, Inches(0.85), Inches(6.62), Inches(11.6), Inches(0.4),
         size=12, color=TEAL, font=SANS)
    if line:
        strip(s, line, Inches(6.75))
    return s


def save(prs, path):
    try:
        prs.save(path)
        return path
    except PermissionError:
        alt = path.replace(".pptx", "_NEW.pptx")
        prs.save(alt)
        return alt
