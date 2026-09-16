# -*- coding: utf-8 -*-
"""Dump S03.pptx for review: every slide's face text and the head of its notes, plus contact sheets.

    python S03_yemen/review_dump.py      # after build.py and series/preview.py

Writes S03_yemen/.build/deck_dump.txt and S03_yemen/.build/sheets/sheet_NN.png (24 slides each), so a
reviewer can read the whole deck as text and look at it in a handful of images instead of 200 PNGs.
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "series"))
from pptx import Presentation                                       # noqa: E402
import preview                                                      # noqa: E402

WORK = os.path.join(HERE, ".build")


def dump():
    p = Presentation(os.path.join(HERE, "S03.pptx"))
    lines = []
    for i, s in enumerate(p.slides, 1):
        hidden = " [HIDDEN]" if s._element.get("show") == "0" else ""
        face = " | ".join(sh.text_frame.text.replace("\n", " / ") for sh in s.shapes
                          if sh.has_text_frame and sh.text_frame.text.strip())
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        head = " ".join(notes.split("\n")[:3])[:220]
        lines.append("### %d%s\nFACE: %s\nNOTES: %s\n" % (i, hidden, face, head))
    path = os.path.join(WORK, "deck_dump.txt")
    os.makedirs(WORK, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return path, len(p.slides)


def sheets(n):
    out = os.path.join(WORK, "sheets")
    os.makedirs(out, exist_ok=True)
    made = []
    for k, start in enumerate(range(1, n + 1, 24), 1):
        path = os.path.join(out, "sheet_%02d.png" % k)
        preview.sheet(os.path.join(HERE, "S03_preview"), path, cols=4, width=420,
                      only=list(range(start, min(start + 24, n + 1))))
        made.append(path)
    return made


if __name__ == "__main__":
    path, n = dump()
    print("%s  %d slides" % (path, n))
    for p in sheets(n):
        print(p)
