# -*- coding: utf-8 -*-
"""The vision gate — every rule in docs/VISION.md that a script can check, checked.

    python tools/check_vision.py S06_hadramawt_bahrayn            # exits non-zero on any breach
    python tools/check_vision.py S06_hadramawt_bahrayn --rule N1  # one rule only
    python tools/check_vision.py S06_hadramawt_bahrayn --slides 1-25

WHY. Daniyal, 2026-10-08: "we need to make the system that generates our content always within the vision
given my feedback." His feedback had been arriving evening by evening and being fixed slide by slide, and the
same complaints came back: notes he could not find his line in, English on a face that meant nothing to the
room, maps that did not move with the story. A rule he has had to repeat is a rule nothing enforces. This
file is where they are enforced. build.py runs it; a deck that breaks a rule is not shown to him.

It reads the BUILT deck, not the intentions behind it: the .pptx, its notes, .build/deck.json (which card is
on which slide, written by build.py), the Map Studio scenes those slides were rendered from, the evening's
timeline.json, and the pool. Rule ids are the ids in docs/VISION.md. To add a rule: write it there with the
words of Daniyal's that produced it, then add its check here — never the fix alone.
"""
import argparse
import json
import os
import re
import sys
import unicodedata

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "series"))
sys.path.insert(0, HERE)

import mapkit                                             # noqa: E402
import notes2                                             # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

SCENES = os.path.join(ROOT, "tools", "mapstudio", "scenes")
SCAFFOLD = "PLACEHOLDER::"
ARABIC = re.compile("[؀-ۿݐ-ݿﭐ-﷿ﹰ-﻿]")
HONORIFIC = re.compile("[ؐ-ؕﷺﷻ]")
TRANSLIT = re.compile("[Ā-ſḀ-ỿʿʾ]")       # ā ī ū ḥ ṣ ḍ ṭ ẓ ʿ ʾ …

# W1 — a Companion is named. A common noun standing where a name belongs.
COMMON_NOUN = re.compile(r"\b(?:one|a|an|the|this|that)\s+(?:man|commander|general|leader|rider|governor)\b"
                         r"(?!\s+(?:of the Book|called|named))", re.I)
# M11 — a key that says something was done and not who did it
PARTICIPLE = re.compile(r"^(?:taken|settled|besieged|opened|closed|drawn in|held|lost|killed|defeated|"
                        r"captured|broken|relieved|crossed|sealed|cut off|surrounded|retaken)\b", re.I)
# W3 — commentary where a rendering belongs
GLOSS_WORDS = re.compile(r"\b(?:the image is|image of|metaphor|literally|that is to say|i\.e\.)", re.I)
# N6 — the production's own apparatus, in the part of the notes he reads at the lectern
APPARATUS = re.compile(r"`|\.md\b|\bDECISIONS\b|(?<![\w/])#\d+\b|§\s*\d|\b[A-Z]{3}/E-[A-Z]+\d+|\bE-[A-Z]{2,3}\d+\b|"
                       r"CONTENT\.md|RUNSHEET|build\.py")
# I2 — an image is a place, an object, an animal, a landscape
DEPICTS = re.compile(r"\b(?:portrait|his face|her face|their faces|depict(?:ing)? (?:the )?(?:Prophet|Companion)|"
                     r"a man's face|close-up of a (?:man|woman))", re.I)

STOP = {"the", "and", "from", "with", "then", "his", "her", "their", "for", "not", "all", "one", "two", "who",
        "right", "left", "front", "last", "first", "road", "army", "men", "side", "other", "still", "come"}
PLAIN_PLACES = {"medina", "mecca", "yemen", "oman", "bahrayn", "arabia", "iraq", "persia", "rome", "syria",
                "kinda", "hajar", "mahra", "najd", "egypt", "jerusalem", "damascus", "basra", "kufa", "badr"}


def norm(s):
    s = HONORIFIC.sub("", s or "")
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9 ]+", "", s.replace("ʿ", "").replace("ʾ", "").lower()).strip()


def has_proper_name(s, lexicon):
    if ARABIC.search(s) or TRANSLIT.search(s):
        return True
    return any(w in lexicon or w in PLAIN_PLACES for w in norm(s).split())


class Report(object):
    def __init__(self, only=None, slides=None):
        self.rows, self.only, self.slides = [], only, slides

    def add(self, rule, where, msg):
        if self.only and rule != self.only:
            return
        if self.slides and isinstance(where, int) and not (self.slides[0] <= where <= self.slides[1]):
            return
        self.rows.append((rule, where, msg))


def load_deck(folder):
    from pptx import Presentation
    stem = os.path.basename(folder.rstrip("/\\")).split("_")[0]
    cands = [os.path.join(folder, stem + x) for x in ("_NEW.pptx", ".pptx")]
    cands = [c for c in cands if os.path.exists(c)]
    if not cands:
        raise SystemExit("no deck in %s — run build.py first" % folder)
    path = max(cands, key=os.path.getmtime)
    return path, Presentation(path)


def slide_facts(s):
    f = {"headline": "", "kicker": "", "speaker": "", "scene": [], "placeholder": False, "brief_on_face": "",
         "boxes": [], "arabic": "", "has_map": False, "names": []}
    for sh in s.shapes:
        nm = sh.name or ""
        txt = sh.text_frame.text.strip() if sh.has_text_frame else ""
        f["names"].append(nm)
        if nm.startswith(SCAFFOLD):
            f["placeholder"] = True
            if nm.endswith("brief"):
                f["brief_on_face"] = txt
            continue
        if nm == "headline":
            f["headline"] = txt
        elif nm == "kicker":
            f["kicker"] = txt
        elif nm == "speaker":
            f["speaker"] = txt
        elif nm == "scene":
            f["scene"].append(txt)
        elif nm == "map" or nm.startswith("map "):
            f["has_map"] = True
        elif nm.startswith("person:") or nm.startswith("child:") or nm.startswith("marriage:"):
            f["tree"] = True
        elif txt:
            f["boxes"].append(txt)
            if ARABIC.search(txt) and len(ARABIC.findall(txt)) > 12 and not f["arabic"]:
                f["arabic"] = txt
    f["notes"] = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
    return f


def lexicon_of(scene_files):
    lex = set()
    for path in scene_files:
        try:
            d = json.load(open(path, encoding="utf-8"))
        except Exception:
            continue
        for o in d.get("objects", []):
            # names only — a settlement, a force, a named region. A caption ("right and left") is not a name,
            # and letting captions in is how the first run passed the title "Right and left".
            if o.get("type") == "label" and str(o.get("note", "")) != "name":
                continue
            for key in ("name", "label") + (("text",) if o.get("type") == "label" else ()):
                for w in norm(o.get(key, "")).split():
                    if len(w) > 2 and w not in STOP:
                        lex.add(w)
    return lex


def check(folder, only=None, slides=None):
    R = Report(only, slides)
    folder = folder.rstrip("/\\")
    path, prs = load_deck(folder)
    build = os.path.join(folder, ".build")
    meta = json.load(open(os.path.join(build, "deck.json"), encoding="utf-8")) \
        if os.path.exists(os.path.join(build, "deck.json")) else {}
    layered = json.load(open(os.path.join(build, "slides.json"), encoding="utf-8")) \
        if os.path.exists(os.path.join(build, "slides.json")) else {}
    kinds = {int(k): v for k, v in (meta.get("slides") or {}).items()}
    lex = lexicon_of([os.path.join(SCENES, f) for f in os.listdir(SCENES) if f.endswith(".json")])

    facts = {i: slide_facts(s) for i, s in enumerate(prs.slides, 1)}

    # ------------------------------------------------------------------ titles and faces
    for i, f in facts.items():
        kind = (kinds.get(i) or {}).get("kind", "")
        face = [f["headline"], f["kicker"], f["speaker"]] + f["scene"] + f["boxes"]
        # F1 — the title and every part divider name a place or a person
        looks_like_divider = not f["headline"] and f["boxes"] and (
            i == 1 or re.match(r"^Part [IVX]+$", f["boxes"][0].strip()))
        if kind in ("title", "section") or (not kind and looks_like_divider):
            body = " ".join(b for b in f["boxes"] if not re.match(r"^(Part [IVX]+|Evening \d+)$", b.strip()))
            body = re.sub(r"\b\d+\s*(?:–|→|-)?\s*\d*\s*AH\b", "", body)
            if body.strip() and not has_proper_name(body, lex):
                R.add("F1", i, "the title %r names no place and no person" % body.strip()[:60])
        # W1 — a Companion is named
        for t in face:
            if t and t[:1] not in '"“' and COMMON_NOUN.search(t):
                R.add("W1", i, "a common noun where a name belongs: %r" % COMMON_NOUN.search(t).group(0))
        if kind == "bridge" and COMMON_NOUN.search(notes2.top_tier(f["notes"])):
            R.add("W1", i, "the bridge's spoken line uses %r — name him"
                  % COMMON_NOUN.search(notes2.top_tier(f["notes"])).group(0))
        # F6 — a kicker says when AND where
        if f["kicker"] and re.fullmatch(r"[\d\s–→,.-]+", f["kicker"]):
            R.add("F6", i, "the kicker is only %r — say the year as a year, and the place" % f["kicker"])
        # Q1 / F4 / W3 — a quotation
        if f["arabic"] and not f.get("tree") and kind not in ("tree",):
            if not f["speaker"]:
                R.add("Q1", i, "a quotation with no speaker on the face (%s)" % (f["headline"] or "no headline"))
            prev = (kinds.get(i - 1) or {}).get("kind", "")
            if prev != "map" and not f["scene"] and kind != "words":
                R.add("F4", i, "a quotation with no scene line and no map before it (%s)" % f["headline"])
            for t in f["boxes"]:
                if not ARABIC.search(t) and GLOSS_WORDS.search(t):
                    R.add("W3", i, "commentary inside the rendering: %r" % GLOSS_WORDS.search(t).group(0))
        # I1 / I2 — an image placeholder carries its brief
        if f["placeholder"]:
            j = f["notes"].find(notes2.BRIEF)
            brief = f["notes"][j + len(notes2.BRIEF):].strip() if j >= 0 else ""
            if notes2.words(brief) < 25:
                R.add("I1", i, "an image placeholder with no paste-ready IMAGE BRIEF in its notes")
            elif DEPICTS.search(brief):
                R.add("I2", i, "the brief asks for %r — places, objects, animals and landscape only"
                      % DEPICTS.search(brief).group(0))

    # ------------------------------------------------------------------ the notes
    for i, f in facts.items():
        notes = f["notes"]
        if notes2.words(notes) < 45:
            continue
        top = notes2.top_tier(notes)
        if not re.search(r"(?m)^(SAY|READ the Arabic|BRIDGE|CHECKPOINT|THE CLOSE|TONIGHT|TREE|BOOKEND)", top):
            R.add("N1", i, "the notes do not open with a SAY block (%d words before any rule)" % notes2.words(top))
            continue
        w = notes2.words(top)
        if w > notes2.TOP_WORDS:
            R.add("N1", i, "%d words above the DETAIL rule; the lectern tier holds %d" % (w, notes2.TOP_WORDS))
        cues = re.findall(r"(?m)^\d+\.\s+(.+)$", top)
        if len(cues) > notes2.CUES:
            R.add("N1", i, "%d cues on one slide; over %d the card wants splitting" % (len(cues), notes2.CUES))
        for c in cues:
            if notes2.words(c) > notes2.CUE_WORDS:
                R.add("N1", i, "a cue of %d words (a cue is a glance): %r" % (notes2.words(c), c[:56] + "…"))
        warns = re.findall(r"(?m)^⚠.*$", top)
        if len(warns) > 2:
            R.add("N4", i, "%d warnings above SAY; two at most, the rest under BACKGROUND" % len(warns))
        for wl in warns:
            if notes2.words(wl) > notes2.WARN_WORDS + 1:
                R.add("N4", i, "a lectern warning of %d words: %r" % (notes2.words(wl), wl[:50] + "…"))
        if re.search(r"(?m)^\d+\..*\n\d+\.", top):
            R.add("N2", i, "cues with no blank line between them")
        m = APPARATUS.search(top)
        if m:
            R.add("N6", i, "production apparatus in the lectern tier: %r" % m.group(0))

    # ------------------------------------------------------------------ N5 — no beat is dropped
    pool_path = os.path.join(ROOT, "L02_baarah_saal", "CONTENT.md")
    cards = meta.get("cards") or []
    if cards and os.path.exists(pool_path):
        import build_full_deck as B
        pool = B.pool_cards()
        txt = "\n".join(open(os.path.join(ROOT, d, "CONTENT.md"), encoding="utf-8").read() for d in B.POOLS)
        for cid in cards:
            m = re.search(r"^### " + re.escape(cid) + r" · .*?(?=^### |\Z)", txt, re.M | re.S)
            bm = re.search(r"\*\*Beats:\*\*[ \t]*\n((?:[ \t]*\d+\.[^\n]*\n?)+)", m.group(0)) if m else None
            raw = len(re.findall(r"(?m)^[ \t]*\d+\.", bm.group(1))) if bm else 0
            if cid in pool and raw != len(pool[cid]["beats"]):
                R.add("N5", cid, "%d beats on the card, %d reach the deck" % (raw, len(pool[cid]["beats"])))
            if cid in pool and not pool[cid]["beats"]:
                R.add("N5", cid, "the card has no beats — nothing for the SAY block")

    # ------------------------------------------------------------------ the maps
    by_scene = {}
    for n, (scene, a, b) in ((int(k), v) for k, v in layered.items()):
        by_scene.setdefault(scene, []).append((a, b, n))
    for scene, uses in sorted(by_scene.items()):
        sp = os.path.join(SCENES, scene + ".json")
        if not os.path.exists(sp):
            R.add("M0", scene, "the scene file is missing")
            continue
        d = json.load(open(sp, encoding="utf-8"))
        for msg in mapkit.check_scene(d) + mapkit.check_crowding(d):
            R.add(msg.split(" ", 1)[0], scene, msg.split(" ", 1)[1])
        for msg in mapkit.check_ranges(d, sorted((a, b) for a, b, _ in uses)):
            R.add("M3", scene, msg.lstrip(": "))
        for fac in d.get("factions", []):
            if fac.get("id") == "f1" and fac.get("color", "").upper() != mapkit.GREEN:
                R.add("M14", scene, "Muslim forces are %s; the one green is %s" % (fac.get("color"), mapkit.GREEN))
        # keys: who did it (M11), and is the place on the map (M12)
        for a, b, n in uses:
            f = facts.get(n)
            if not f:
                continue
            on_map = set()
            for o in d.get("objects", []):
                if o.get("step", 1) <= b:
                    for key in ("name", "text", "label"):
                        if o.get(key):
                            on_map.add(norm(o[key]))
            pairs = list(zip(f["boxes"][0::2], f["boxes"][1::2]))
            for label, value in pairs:
                if PARTICIPLE.match(value.strip()) and " by " not in " " + value + " " \
                        and not HONORIFIC.search(label + value):
                    R.add("M11", n, "the key %r — %r does not say who did it" % (label, value))
                for part in re.split(r",| and ", label):
                    p = norm(re.sub(r"\(.*?\)", "", part))
                    if not p or "(" in part:
                        continue
                    if not any(p in m or m in p for m in on_map if m):
                        R.add("M12", n, "the key names %r, which is not on this map and is not glossed"
                              % part.strip())

    # M1 — a part opens on a map (or a family tree), unless the runsheet exempts it
    exempt = set(meta.get("no_map_parts") or [])
    order = sorted(kinds)
    for idx, i in enumerate(order):
        if kinds[i].get("kind") != "section":
            continue
        part = kinds[i].get("part", "")
        nxt = [kinds[j].get("kind") for j in order[idx + 1:idx + 4]]
        first = next((k for k in nxt if k not in ("bridge",)), "")
        if first not in ("map", "tree") and part not in exempt:
            R.add("M1", i, "%s opens on a %s slide, not on a map of the situation" % (part or "this part", first or "?"))

    # D1 — a moment of faith has a slide of its own
    for cid in meta.get("iman") or []:
        hits = [i for i, k in kinds.items() if k.get("card") == cid and facts[i]["arabic"]]
        if not hits:
            R.add("D1", cid, "marked as a moment of faith, and no slide carries its Arabic")

    # ------------------------------------------------------------------ the Line
    tl = os.path.join(folder, "timeline.json")
    if os.path.exists(tl):
        t = json.load(open(tl, encoding="utf-8"))
        if len(t.get("lanes", [])) > 2:
            R.add("L1", "timeline", "%d strands on the Line; one or two" % len(t["lanes"]))
        if len(t.get("events", [])) > 8:
            R.add("L1", "timeline", "%d events on the Line; eight at most" % len(t["events"]))

    # I1 — one file with every brief
    if any(f["placeholder"] for f in facts.values()) and not os.path.exists(os.path.join(folder, "IMAGE_BRIEFS.md")):
        R.add("I1", "IMAGE_BRIEFS.md", "the evening has image placeholders and no IMAGE_BRIEFS.md")
    return path, R


def main():
    ap = argparse.ArgumentParser(description="Fail a deck that breaks a rule in docs/VISION.md.")
    ap.add_argument("folder")
    ap.add_argument("--rule")
    ap.add_argument("--slides", help="e.g. 1-25")
    a = ap.parse_args()
    sl = tuple(int(x) for x in a.slides.split("-")) if a.slides else None
    path, R = check(a.folder, a.rule, sl)
    print("%s" % os.path.relpath(path, ROOT))
    if not R.rows:
        print("   vision: every checkable rule holds")
        return 0
    last = None
    for rule, where, msg in sorted(R.rows, key=lambda r: (r[0], str(r[1]).zfill(4))):
        if rule != last:
            print("\n %s" % rule)
            last = rule
        print("   %-12s %s" % (("slide %d" % where) if isinstance(where, int) else where, msg))
    by = {}
    for rule, _, _ in R.rows:
        by[rule] = by.get(rule, 0) + 1
    print("\n%d breach(es) of %d rule(s): %s" % (len(R.rows), len(by), "  ".join("%s=%d" % kv for kv in sorted(by.items()))))
    return 1


if __name__ == "__main__":
    sys.exit(main())
