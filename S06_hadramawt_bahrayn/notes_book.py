# -*- coding: utf-8 -*-
"""S06_notes.pdf and S06_notes.md — every slide of S06.pptx with its full speaker notes, for reading at home.

    python S06_hadramawt_bahrayn/notes_book.py          # after build.py

WHY (DECISIONS.md #60). Daniyal delivers from the deck and its notes only. So the home reading is the exact
thing he will deliver from: each slide as the room will see it, and under it every word of its notes. On a
moving map, the picture of each click sits under its "▶ CLICK" line, so the build-up can be rehearsed away
from PowerPoint. The PDF is A5, to be read on a phone; the .md is the same text, for searching and for copying.

The page-making is evening 5's (S05_yamama_dead_oman_mahra/notes_book.py), imported, not copied. Two things
are this evening's own:

  * THE NOTES ARE TWO TIERS (series/notes2.py): the labels that open a paragraph are SAY, LESSON, ASK and
    the drawn rules — DETAIL, BACKGROUND, IMAGE BRIEF.
  * A MAP SLIDE'S FACE SHOWS ITS LAST FRAME. PowerPoint exports a slide with every click's layer switched on
    at once — every caption of the slide on top of every other. That is not a picture the room ever sees, and
    Daniyal read it as "a lot of rush of labels" (2026-10-08). So the map on the face is replaced by the
    rendered frame the slide ENDS on; the frames in between sit under their clicks.
"""
import html
import importlib.util
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
_spec = importlib.util.spec_from_file_location(
    "s05_notes", os.path.join(ROOT, "S05_yamama_dead_oman_mahra", "notes_book.py"))
N5 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(N5)

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(errors="replace")

DECK = os.path.join(HERE, "S06.pptx")
MAPS = os.path.join(HERE, "visuals", "maps")
BUILD = os.path.join(HERE, ".build")
WORK = os.path.join(BUILD, "notes")            # the printed page and the pasted faces: regenerable, not kept
RULE = re.compile(r"^─{4,}\s*(.*?)\s*─{4,}$")
N5.HEAD = re.compile(r"^(SAY\b|LESSON —|ASK —|READ the Arabic|WHO IS SPEAKING|BRIDGE|CHECKPOINT|THE CLOSE|BOOKEND|"
                     r"TONIGHT|TREE|IF YOU |▶ NEXT SLIDE|⚠|─{4,})")
N5.ITEM = re.compile(r"^\s*(\d+\.\s|▶ |— |LESSON —|ASK —|SAY\b|READ the Arabic|WHO IS SPEAKING|IF YOU |⚠|─{4,}|"
                     r"BRIDGE|CHECKPOINT|THE CLOSE|BOOKEND|TONIGHT|Each part)")
TITLE = "Evening 6 — Bahrayn and Ḥaḍramawt, and the other side of the desert"
LEAD = ("Every slide as the room will see it, and under it every word of its notes. On a moving map the slide "
        "shows the frame it ends on, and each click's picture sits under its ▶ CLICK line. The short lines at the "
        "top are what you say; everything under a drawn rule is there if you need it.")


def last_frame(face, slide, scene, size):
    """The exported face of a map slide, with the map replaced by the frame the slide ends on."""
    from PIL import Image
    step = os.path.join(MAPS, "%s-step-%02d.png" % (scene[0], scene[2]))
    pics = [sh for sh in slide.shapes if sh.shape_type == 13 and (sh.name or "").lower().startswith("map")]
    if not os.path.exists(step) or not pics:
        return N5.jpeg(face)
    big = max(pics, key=lambda sh: sh.width * sh.height)
    im = Image.open(face).convert("RGB")
    fx, fy = im.width / float(size[0]), im.height / float(size[1])
    box = (int(big.left * fx), int(big.top * fy), int((big.left + big.width) * fx), int((big.top + big.height) * fy))
    frame = Image.open(step).convert("RGB").resize((box[2] - box[0], box[3] - box[1]), Image.LANCZOS)
    im.paste(frame, box[:2])
    out = os.path.join(WORK, "face-%03d.jpg" % int(re.search(r"(\d+)\.png$", face, re.I).group(1)))
    im.save(out, "JPEG", quality=85, optimize=True)
    return out


def md_notes(notes):
    """The notes as Markdown: one line a line, a rule as a bold label, Arabic in a blockquote of its own (#25)."""
    out = []
    # a quotation folded into running prose by the notes ("… on the same page: > وقد قيل …") gets its own line back
    notes = re.sub(r"\s+>\s+(?=[؀-ۿ—*]|English:)", "\n> ", notes)
    for raw in notes.replace("⁨", "").replace("⁩", "").split("\n"):
        line = raw.rstrip()
        m = RULE.match(line.strip())
        if not line.strip():
            out.append("")
        elif m:
            out += ["", "---", "", "**%s**" % m.group(1), ""]
        elif line.lstrip().startswith(">"):
            out += ["", "> " + line.lstrip("> ").strip(), ""]
        elif N5.AR.search(line) and not N5.LATIN.search(line):
            out += ["", "> " + line.strip(), ""]
        else:
            out.append(line + "  ")
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip()


def build():
    from pptx import Presentation
    prs = Presentation(DECK)
    slides = list(prs.slides)
    size = (prs.slide_width, prs.slide_height)
    layered = json.load(open(os.path.join(BUILD, "slides.json"), encoding="utf-8"))
    os.makedirs(WORK, exist_ok=True)
    imgdir = N5.preview.export(DECK)
    faces = {}
    for f in os.listdir(imgdir):
        m = re.search(r"(\d+)\.png$", f, re.I)
        if m:
            faces[int(m.group(1))] = os.path.join(imgdir, f)

    page = ['<!DOCTYPE html><html><head><meta charset="utf-8"><title>%s — the notes book</title>'
            '<!--EMBED-FONTS--><style>%s</style></head><body>' % (html.escape(TITLE), N5.CSS),
            '<div><h1>%s</h1><p class="lead">%s</p></div>' % (html.escape(TITLE), html.escape(LEAD))]
    md = ["# %s — the notes book" % TITLE, "", LEAD, "",
          "*Made by `S06_hadramawt_bahrayn/notes_book.py` from `S06.pptx`. Do not edit: rebuild the deck, then this.*", ""]
    for n, s in enumerate(slides, 1):
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        scene = layered.get(str(n))
        body = []
        for cls, block in N5.paragraphs(notes):
            if RULE.match(block.strip()):
                cls = "head"
            body.append(N5.para(block, cls))
            m = N5.CLICK.match(block)
            if m and scene:
                png = os.path.join(MAPS, "%s-step-%02d.png" % (scene[0], scene[1] + int(m.group(1))))
                if os.path.exists(png):
                    body.append('<img class="step" src="%s">' % N5.url(N5.jpeg(png, width=960)))
        face = faces.get(n)
        shown = (last_frame(face, s, scene, size) if scene else N5.jpeg(face)) if face else None
        page.append('<div class="slide"><div class="no">SLIDE %d</div>%s<div class="notes">%s</div></div>'
                    % (n, '<img class="face" src="%s">' % N5.url(shown) if shown else "", "".join(body)))

        named = {(sh.name or ""): sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame}
        head = named.get("headline") or next((v for v in named.values() if v), "")
        md += ["", "## Slide %d — %s" % (n, " ".join(head.split())), ""]
        if named.get("kicker"):
            md += ["*%s*" % named["kicker"], ""]
        if scene:
            md += ["*A map: `%s`, steps %d to %d — %d click(s).*" % (scene[0], scene[1], scene[2], scene[2] - scene[1]), ""]
        md += [md_notes(notes) if notes.strip() else "*(no notes)*"]
    page.append("</body></html>")

    src = os.path.join(WORK, "S06_notes.html")
    with open(src, "w", encoding="utf-8") as f:
        f.write("".join(page))
    pdf = N5.embed_fonts.to_pdf(src)
    final = os.path.join(HERE, "S06_notes.pdf")
    if os.path.exists(final):
        os.remove(final)
    os.replace(pdf, final)
    with open(os.path.join(HERE, "S06_notes.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(md).rstrip() + "\n")
    print("%-24s %3d slides  %8.1f KB" % (os.path.basename(final), len(slides), os.path.getsize(final) / 1024))
    print("%-24s %3d slides  %8.1f KB" % ("S06_notes.md", len(slides), os.path.getsize(os.path.join(HERE, "S06_notes.md")) / 1024))


if __name__ == "__main__":
    build()
