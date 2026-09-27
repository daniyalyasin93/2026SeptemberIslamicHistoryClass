# -*- coding: utf-8 -*-
"""S05_notes.pdf — every slide of S05.pptx with its full speaker notes beneath it, for reading at home.

    python S05_yamama_dead_oman_mahra/notes_book.py          # after build.py

WHY (DECISIONS.md #60). Daniyal delivers from the deck and its notes only; the 122-page briefing and the
cue sheet were never opened. So the home reading is the exact thing he will deliver from: each slide as
the room will see it, and under it every word of its notes. On a moving map, the picture of each click
sits under its "▶ CLICK" line, so the build-up can be rehearsed away from PowerPoint.

WHY NOT POWERPOINT'S OWN "NOTES PAGES" EXPORT. It prints a fixed-size notes box and cuts long notes off
at its edge — and these notes are long by design. This is HTML → headless Chrome → PDF, the repo's own
printable path (series/embed_fonts.py), A5 so it reads on a phone, Arabic set in Traditional Arabic.
"""
import html
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "series"))
import embed_fonts                                                   # noqa: E402
import preview                                                       # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(errors="replace")

DECK = os.path.join(HERE, "S05.pptx")
MAPS = os.path.join(HERE, "visuals", "maps")
BUILD = os.path.join(HERE, ".build")
AR = re.compile("[؀-ۿﭐ-﷿ﹰ-﻿]")
LATIN = re.compile("[A-Za-z]")
HEAD = re.compile(r"^(SAY —|QUOTE —|عبرت —|HANDS UP|BACKGROUND|NOTE —|⏱|CARRYING ON|OUT OF TIME|BRIDGE|"
                  r"CHECKPOINT|THE CLOSE|BOOKEND|TONIGHT|⚠)")
CLICK = re.compile(r"^▶ CLICK (\d+)")

CSS = """
@page { size: A5; margin: 9mm 9mm 11mm 9mm; }
body { font-family: 'Segoe UI', sans-serif; font-size: 9.6pt; color: #1C2321; line-height: 1.38; }
.slide { page-break-before: always; }
.slide:first-of-type { page-break-before: auto; }
.no { font: bold 8pt 'Segoe UI'; color: #C49A45; letter-spacing: .06em; margin-bottom: 2mm; }
img.face { width: 100%; border: 0.3mm solid #D8D2C4; }
.notes { margin-top: 3mm; }
.notes p { margin: 0 0 1.6mm 0; white-space: pre-wrap; }
.notes p.head { font-weight: 700; color: #104A43; }
.notes p.ar { direction: rtl; text-align: right; font-family: 'Traditional Arabic', serif; font-size: 15pt;
              line-height: 1.5; }
.notes p.click { font-weight: 700; color: #7A2E2E; margin-top: 2.5mm; }
span.arr { font-family: 'Traditional Arabic', serif; font-size: 13pt; unicode-bidi: isolate; }
img.step { width: 72%; display: block; margin: 1mm 0 3mm 4mm; border: 0.3mm solid #D8D2C4; }
h1 { font: bold 16pt Georgia; color: #104A43; margin: 0 0 2mm 0; }
.lead { color: #5A6862; margin-bottom: 6mm; }
"""


def jpeg(png, width=None):
    from PIL import Image
    out = os.path.splitext(png)[0] + (".book.jpg" if width else ".jpg")
    if not os.path.exists(out) or os.path.getmtime(out) < os.path.getmtime(png):
        im = Image.open(png).convert("RGB")
        if width and im.width > width:
            im = im.resize((width, int(im.height * width / im.width)))
        im.save(out, "JPEG", quality=85, optimize=True)
    return out


def url(path):
    return "file:///" + os.path.abspath(path).replace("\\", "/")


ARRUN = re.compile("[\u0600-\u06FF\uFB50-\uFDFF\uFE70-\uFEFF](?:[\u0600-\u06FF\uFB50-\uFDFF\uFE70-\uFEFF"
                   "\\s\u060C\u061B\u061F\u00AB\u00BB:.,!()\\-\u2013\u2014\\d]*[\u0600-\u06FF\uFB50-\uFDFF\uFE70-\uFEFF])?")
MD = re.compile(r"\*\*|(?<!\w)\*(?!\s)|(?<=\S)\*(?!\w)|`")
ITEM = re.compile(r"^\s*(\d+\.\s|\^|▶ CLICK|— |Then |CARRYING ON|OUT OF TIME|⏱|SAY —|NOTE —|⚠|BACKGROUND|"
                  r"HANDS UP|عبرت —|QUOTE —|CHECKPOINT|BRIDGE|THE CLOSE|BOOKEND|TONIGHT|\s{2,}\d+\.)")


def mixed(text):
    """English prose with Arabic inside it: the Arabic runs are isolated right-to-left and the rest reads left
    to right — tools/bidi_fix.py's fix, for the page instead of the editor."""
    out, last = [], 0
    for m in ARRUN.finditer(text):
        out.append(html.escape(text[last:m.start()]))
        out.append('<span class="arr" dir="rtl">%s</span>' % html.escape(m.group(0)))
        last = m.end()
    out.append(html.escape(text[last:]))
    return "".join(out)


def para(line, cls=""):
    line = MD.sub("", line.replace("\u2068", "").replace("\u2069", "")).rstrip()
    if not line.strip():
        return ""
    ar, lat = len(AR.findall(line)), len(LATIN.findall(line))
    if lat == 0 and ar:
        return '<p class="ar" dir="rtl">%s</p>' % html.escape(line)
    if CLICK.match(line):
        cls = "click"
    elif HEAD.match(line):
        cls = cls or "head"
    return "<p%s>%s</p>" % (' class="%s"' % cls if cls else "", mixed(line))


def paragraphs(notes):
    """The notes as paragraphs: a hard-wrapped run of lines is one paragraph; a list item, a label, a click
    or a quotation line starts a new one."""
    out = []
    for block in notes.split("\n\n"):
        lines = block.split("\n")
        if lines and lines[0].startswith("QUOTE — "):
            out.append(("head", "QUOTE —"))
            out.append(("", lines[0][8:]))
            out.extend(("", l) for l in lines[1:])
            continue
        cur = ""
        for l in lines:
            if cur and (ITEM.match(l) or not l.strip()):
                out.append(("", cur))
                cur = l
            else:
                cur = (cur + " " + l.strip()) if cur else l
        if cur:
            out.append(("", cur))
    return out


def build():
    from pptx import Presentation
    slides = list(Presentation(DECK).slides)
    layered = json.load(open(os.path.join(BUILD, "slides.json"), encoding="utf-8"))
    imgdir = preview.export(DECK)
    faces = {}
    for f in os.listdir(imgdir):
        m = re.search(r"(\d+)\.png$", f, re.I)
        if m:
            faces[int(m.group(1))] = os.path.join(imgdir, f)

    out = ['<!DOCTYPE html><html><head><meta charset="utf-8"><title>Evening 5 — the notes book</title>'
           '<!--EMBED-FONTS--><style>%s</style></head><body>' % CSS,
           '<div><h1>Evening 5 — what al-Yamāma cost</h1><p class="lead">Every slide as the room will see it, and '
           'under it every word of its notes. On a moving map, each click\'s picture sits under its ▶ CLICK line. '
           'Read it twice; deliver from the deck.</p></div>']
    for n, s in enumerate(slides, 1):
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        scene = layered.get(str(n))
        body = []
        for cls, block in paragraphs(notes):
            body.append(para(block, cls))
            m = CLICK.match(block)
            if m and scene:
                step = scene[1] + int(m.group(1))
                png = os.path.join(MAPS, "%s-step-%02d.png" % (scene[0], step))
                if os.path.exists(png):
                    body.append('<img class="step" src="%s">' % url(jpeg(png, width=960)))
        face = faces.get(n)
        out.append('<div class="slide"><div class="no">SLIDE %d</div>%s<div class="notes">%s</div></div>'
                   % (n, '<img class="face" src="%s">' % url(jpeg(face)) if face else "", "".join(body)))
    out.append("</body></html>")
    os.makedirs(BUILD, exist_ok=True)
    page = os.path.join(BUILD, "S05_notes.html")
    with open(page, "w", encoding="utf-8") as f:
        f.write("".join(out))
    pdf = embed_fonts.to_pdf(page)
    final = os.path.join(HERE, "S05_notes.pdf")
    if os.path.exists(final):
        os.remove(final)
    os.replace(pdf, final)
    print("%-24s %3d slides  %8.1f KB" % (os.path.basename(final), len(slides), os.path.getsize(final) / 1024))


if __name__ == "__main__":
    build()
