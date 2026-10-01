# -*- coding: utf-8 -*-
"""Build S05.pptx — evening 5: what al-Yamāma cost, and ʿIkrima's ؓ road to the last front.

    python S05_yamama_dead_oman_mahra/make_timeline.py   # -> visuals/line_s05_*.png      (the Line)
    python S05_yamama_dead_oman_mahra/make_maps.py       # -> scenes + visuals/maps/…      (maps and their click layers)
    python S05_yamama_dead_oman_mahra/build.py           # -> S05.pptx + S05.pdf
    python S05_yamama_dead_oman_mahra/notes_book.py      # -> S05_notes.pdf                (every slide with its notes)

WHAT IS NEW ON EVENING 5 (DECISIONS.md #56–#61, and the second pass #62–#66).
  * MAPS BUILD UP ON CLICKS (#58), AND A MOVING CARD'S MAP IS ITS SLIDE (#63). The first build put a map in
    front of the card and then told the card again on its own slide; Daniyal saw the story told twice. Now
    the map carries the card: its beats are in the map's notes with "▶ CLICK n" on the beat the map moves
    on, and a card with an Arabic quotation keeps one more slide after the map, for the words only.
  * WHAT A CLICK HAS FINISHED WITH FADES (#64) — make_maps.py sets it; series/anim.py plays it.
  * FAMILY TREES (#65). Each house opens on one diagram slide. The face carries names and lines only; the
    notes give every relative one sourced sentence, and the BACKGROUND keeps the beats of the cards the
    second pass moved off the running order — nothing researched is lost.
  * CHECKPOINTS, NOT MID-EVENING CLOSES (#59): three, and the close. «Tonight» and «Next week» only at the close.
  * THE NOTES ARE THE LECTERN (#60): warnings first, then SAY, then the quotation, the lesson, the hands-up;
    background last. No file paths or build notes anywhere in the pane.

The face rules for the cards evening 4 built behind its close are carried over from
S04_kinda_butah_yamama/build.py — loaded from its file — so a card reads the same on both evenings.
"""
import importlib.util
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "series"))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import deck2 as D                                                   # noqa: E402
import build_full_deck as B                                         # noqa: E402
import anim                                                         # noqa: E402
from pptx.util import Inches, Pt                                    # noqa: E402
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR                     # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR                          # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE                       # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(errors="replace")

_spec = importlib.util.spec_from_file_location("s04_build", os.path.join(ROOT, "S04_kinda_butah_yamama", "build.py"))
S4 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(S4)

VIS = os.path.join(HERE, "visuals")
# short citation forms the runsheet uses, so an inline citation turned to English keeps its book
for _pair in (("سیر الراشدون", "Siyar (al-Rashidun)"), ("سير الخلفاء الراشدين", "al-Rashidun"),
              ("البدایہ", "al-Bidaya"), ("سیر", "Siyar")):
    if _pair not in B.TRANSLIT:
        B.TRANSLIT.append(_pair)
MAPS = os.path.join(VIS, "maps")
BOOKEND = os.path.join(VIS, "bookend")

# ---------------------------------------------------------------------------------------------- faces
FACE_TITLE = dict(S4.FACE_TITLE, **{
    "RCT/E-RC69": "Musaylima's words, at Medina",
    "RCT/E-RC29": "Dabā: the relief from inside",
    "RCT/E-RC30": "Mahra: won with a letter",
    "RCT/E-RC70": "Oman asked for help",
    "ABU/E-Q7": "Zayd b. Thābit ؓ, as a boy",
    "ABU/E-Q3": "What it was gathered from",
    "TSY/E-YK19": "al-Ashʿath b. Qays comes into it",
    # #68 (2026-10-01): the dead of al-Yamāma — Parts II–IV — are headed by the Companion's name, with a plain
    # word or two only where one man has several slides. Daniyal delivers in Urdu; an English sentence on the
    # face is no use to him there, a name is.
    "THO/E-HS14": "al-Barāʾ b. Mālik ؓ",
    "THO/E-HS1": "Abū Ṭalḥa ؓ and Umm Sulaym ؓ",
    "THO/E-HS9": "Umm Sulaym ؓ",
    "THO/E-HS15": "al-Barāʾ ؓ and Anas ؓ",
    "THO/E-HS17": "al-Barāʾ ؓ at Tustar",
    "ZIA/E-ZY12": "Abū ʿAqīl ؓ",
    "ZIA/E-ZY15": "Ḥabīb b. Zayd ؓ, and Umm ʿUmāra ؓ",
    "ZIA/E-ZY18": "ʿAbd Allāh b. Zayd ؓ",
    "AHA/E-AS13": "Zayd b. al-Khaṭṭāb ؓ — the banner",
    "ZIA/E-ZY4": "ʿUmar ؓ, and the man who killed Zayd ؓ",
    "ZIA/E-ZY8": "Zayd ؓ and Maʿn b. ʿAdī ؓ",
    "ZIA/E-ZY9": "Maʿn b. ʿAdī ؓ",
    "ZIA/E-ZY10": "Thābit b. Qays ؓ",
    "AHA/E-AS02": "Abū Ḥudhayfa b. ʿUtba ؓ",
    "AHA/E-AS06": "Sālim ؓ — from Iṣṭakhr",
    "AHA/E-AS07": "Sālim ؓ at Qubāʾ",
    "AHA/E-AS04": "Abū Ḥudhayfa ؓ at Badr",
    "AHA/E-AS09": "Sālim ؓ — one of the four",
    "AHA/E-AS15": "Sālim ؓ — the banner",
    "AHA/E-AS16": "Sālim ؓ and Abū Ḥudhayfa ؓ",
    "AHA/E-AS18": "ʿUmar ؓ on Sālim ؓ",
})
FACE_WHEN = dict(S4.FACE_WHEN, **{
    "RCT/E-RC67": "12 AH, the forts of al-Yamāma",
    "RCT/E-RC68": "12 AH, al-Yamāma",
    "RCT/E-RC69": "12 AH, Medina",
    "RCT/E-RC70": "11–12 AH, Oman",
    "RCT/E-RC29": "11–12 AH, Oman",
    "RCT/E-RC30": "11–12 AH, Mahra",
    "TMW/E-TRN6": "8 AH, the Conquest of Mecca",
    "TMW/E-TRN7": "8 AH, the Conquest of Mecca",
    "TMW/E-TRN8": "8 AH",
    "RCT/E-RC23": "12 AH, Medina",
    "ABU/E-Q7": "1 AH, Medina",
    "ABU/E-Q3": "12 AH, Medina",
    "ABU/E-Q4": "ʿAlī ؓ, on the collection",
    "ABU/E-Q5": "Years later, Medina",
    "KTK/E-KD05": "10 AH, Medina",
    "TSY/E-YK07": "11–12 AH, Ḥaḍramawt",
    "RCT/E-RC33": "11–12 AH, Ḥaḍramawt",
    "TSY/E-YK08": "11–12 AH, Ḥaḍramawt",
    "TSY/E-YK09": "11–12 AH, Ḥaḍramawt",
    "TSY/E-YK10": "11–12 AH, Ḥaḍramawt",
    "TSY/E-YK19": "11–12 AH, Ḥaḍramawt",
    "TSY/E-YK11": "11–12 AH, Ḥaḍramawt",
    "TSY/E-YK12": "11 or 12 AH, al-Nujayr",
    "TSY/E-YK14": "12 AH, Medina",
    "POT/E-PG40": "12 AH, Medina",
    "TSY/E-YK16": "Years later: al-Qādisiyya, Nahāwand",
    "POT/E-PG41": "Years later",
})
FACE_QUOTE = dict(S4.FACE_QUOTE, **{
    # the terms, cut before "and they settled on that" so the list is whole on the face
    "RCT/E-RC67": ("… صالحه على الصفراء والبيضاء والحلقة والكراع، وعلى نصف الرقيق، وعلى حائط من كل قرية",
                   "…made terms with him on the gold and the silver, the armour and the horses, half of the "
                   "slaves, and a walled garden from every village"),
    "RCT/E-RC68": ("ودعاهم خالد إلى الإسلام فأسلموا عن آخرهم ورجعوا إلى الحق",
                   "And Khālid called them to Islam, and they became Muslims to the last of them and returned "
                   "to the truth"),
    # the storm and the vow; the whole covenant is in the notes
    "TMW/E-TRN6": ("وَاللَّهِ لَئِنْ لَمْ يُنْجِ فِي الْبَحْرِ إِلَّا الْإِخْلَاصُ فَإِنَّهُ لَا يُنْجِي فِي الْبَرِّ غَيْرُهُ … أَنْ آتِيَ مُحَمَّدًا حَتَّى أَضَعَ يَدِي فِي يَدِهِ",
                   "By God, if nothing saves at sea but sincerity, then nothing else saves on land either … I "
                   "will go to Muhammad and put my hand in his hand"),
})
FACE_CITE = dict(S4.FACE_CITE, **{
    # the note's citation line carries stray words ("verse as quoted at … the reference"), which reached the face
    "THO/E-HS5": "البدایہ والنہایہ ج۴ ص۲۰۳",
})
# Four hands-up in the notes (ZY4, ZY17, Q3, and TRN6 if the evening gets there); every other card's own
# hands-up line is a spare, named in the runsheet's note instead of prompted at the foot of the pane.
NO_HANDS = set(S4.NO_HANDS) | {
    "RCT/E-RC67", "RCT/E-RC37", "AHA/E-AS02", "AHA/E-AS06", "AHA/E-AS15", "RCT/E-RC29", "KTK/E-KD05",
    "TSY/E-YK07", "RCT/E-RC33", "TSY/E-YK09", "TSY/E-YK14", "TSY/E-YK16"}


def check_face_quotes(pool):
    S4.FACE_QUOTE = FACE_QUOTE          # S4's checker reads its own module global
    S4.check_face_quotes(pool)


# ---------------------------------------------------------------------------------------------- the recap
# The one map that is not a card's own: last week, in one minute, before RC67 (#1 in the runsheet).
RECAP_BEFORE = {
    "RCT/E-RC67": ("s05-01-aqraba", 1, 6, "The day at ʿAqrabāʾ", "Last week, in one minute",
                   [("The garden", "the gate shut, then opened"), ("Musaylima", "falls"),
                    ("The forts", "where the terms were made")],
                   "Last week ended with al-Yamāma taken. Here is the end of that day again, in five moves — "
                   "and then what was signed.",
                   ["Banū Ḥanīfa break, and fall back into a walled garden. The gate shuts behind them.",
                    "al-Barāʾ b. Mālik ؓ: throw me over — over the wall, on a shield raised on spears.",
                    "He fights at the gate and opens it, and the army pours in.",
                    "In a gap in the wall, Musaylima falls.",
                    "The survivors reach the forts. And this is where tonight begins: what was signed there."]),
}

# ---------------------------------------------------------------------------------------------- map = the card's slide
# card id -> (scene, first step, last step, keys, [(beat, what the click lands on), ...]) — #63.
# The beat numbers are the card's own; a click may not sit on a beat the runsheet skips (the build checks).
MAP_FOR = {
    "THO/E-HS17": ("s05-02-tustar", 1, 2,
                   [("al-Yamāma", "the gate, 12 AH"), ("Tustar", "years later — a war not yet reached")],
                   [(1, "Tustar, years later — in a war we have not reached.")]),
    "AHA/E-AS06": ("s05-03-salim", 1, 2,
                   [("Mecca", "Abū Ḥudhayfa's ؓ house"), ("Iṣṭakhr", "in Persia — where Sālim ؓ came from")],
                   [(1, "Iṣṭakhr, in Persia: where he came from. The books give his origin, not his road.")]),
    "TMW/E-TRN6": ("s05-04-ikrima-flight", 1, 2,
                   [("Mecca", "the Conquest, 8 AH"), ("The coast", "a ship, and a storm")],
                   [(3, "He runs for the coast, and takes a ship.")]),
    "TMW/E-TRN7": ("s05-04-ikrima-flight", 2, 4,
                   [("Umm Ḥakīm ؓ", "his wife, with the safe-conduct"), ("Mecca", "she brings him back")],
                   [(3, "She goes after him, down the coast road."),
                    (4, "She reaches him at the shore — and brings him back.")]),
    "RCT/E-RC70": ("s05-05-oman", 1, 6,
                   [("Oman", "Laqīṭ, 'the one with the crown'"), ("Jayfar", "asks Medina for help"),
                    ("ʿIkrima ؓ", "joins them from al-Yamāma")],
                   [(2, "The two lawful rulers, driven to the edges of their own country."),
                    (3, "Jayfar writes to Abū Bakr ؓ, and asks for help."),
                    (4, "Two commanders come from Medina."),
                    (5, "And ʿIkrima's ؓ line from al-Yamāma — they join at Rijām, near Oman."),
                    (6, "Letters before swords: to the rebel's own chiefs.")]),
    "RCT/E-RC29": ("s05-06-daba", 1, 5,
                   [("Dabā", "the great market town"), ("Ṣuḥār", "the lawful rulers"),
                    ("The relief", "from inside the region")],
                   [(3, "The Muslims advance from Ṣuḥār."),
                    (4, "Laqīṭ gets the upper hand, and the line gives way. (Spare hands-up before the next "
                        "click: the column from Medina is losing — who is close enough to help?)"),
                    (5, "The relief: Banū Nājiya and ʿAbd al-Qays — men of that region, not an army from Medina."),
                    (6, "The rebels broken.")]),
    "RCT/E-RC30": ("s05-07-mahra", 1, 5,
                   [("Mahra", "two chiefs at odds"), ("The letter", "half the enemy comes over"),
                    ("The fifth", "to Medina")],
                   [(1, "ʿIkrima ؓ comes into Mahra."),
                    (4, "The letter — and Shikhrīt comes over."),
                    (5, "al-Muṣabbaḥ refuses: the fighting is harder than Dabā."),
                    (6, "And the fifth goes to Medina, carried by Shikhrīt himself.")]),
    "TSY/E-YK07": ("s05-10-kinda", 1, 3,
                   [("Ḥaḍramawt", "Ziyād b. Labīd ؓ, the Prophet's ﷺ governor"), ("Kinda", "its neighbour")],
                   [(2, "The arrangement: ṣadaqa carried both ways."),
                    (6, "'Carry it yourselves' — and it was not carried.")]),
    "TSY/E-YK08": ("s05-10-kinda", 3, 6,
                   [("Ziyād's ؓ camp", "Ḥaḍramawt and al-Sakūn"), ("Banū Muʿāwiya", "of Kinda")],
                   [(5, "Two great camps face each other."),
                    (7, "A strike by night."),
                    (8, "Kinda roused: no ṣadaqa.")]),
    "TSY/E-YK09": ("s05-10-kinda", 6, 8,
                   [("The pastures", "each chief's, held as a stronghold"), ("Shuraḥbīl", "and his son, to Ziyād ؓ")],
                   [(1, "Each chief takes his own pasture as a stronghold."),
                    (4, "One man and his son walk out, to Ziyād's ؓ camp.")]),
    "TSY/E-YK10": ("s05-10-kinda", 8, 10,
                   [("Banū ʿAmr", "the strongest of Kinda"), ("The four kings", "killed at their fires")],
                   [(1, "They ring the pastures in the dark — from five sides."),
                    (4, "The four kings are killed.")]),
    "TSY/E-YK19": ("s05-10-kinda", 10, 13,
                   [("The column", "the property and the captives"), ("al-Ashʿath b. Qays", "of the other branch")],
                   [(1, "The column turns for home."),
                    (2, "Its road passes one man: al-Ashʿath b. Qays."),
                    (5, "He takes the captives back.")]),
    "TSY/E-YK11": ("s05-10-kinda", 13, 16,
                   [("al-Muhājir ؓ", "from Ṣanʿāʾ, ahead of his army"), ("ʿIkrima ؓ", "with the main body"),
                    ("al-Nujayr", "the fort")],
                   [(2, "al-Muhājir ؓ rides ahead with the fastest men."),
                    (3, "Maḥjar al-Zurqān: Kinda breaks, and runs."),
                    (4, "Into al-Nujayr — and ʿIkrima ؓ brings up the main body. The ring closes.")]),
}

# ---------------------------------------------------------------------------------------------- bridges
BRIDGE_BEFORE = {
    "AHA/E-AS13": ("The banner", "12 AH, at al-Yamāma", [
        ("Last week", "Zayd b. al-Khaṭṭāb ؓ at the line, and his vow of silence"),
        ("Now", "what happened to the banner he held"),
        ("And after", "the men who carried the Qurʾān")],
        "Last week you saw the banner at the line, and Zayd b. al-Khaṭṭāb ؓ's vow not to speak until it was "
        "over. Here is what happened to that banner."),
    "AHA/E-AS09": ("Back to the field", "12 AH", [
        ("That was the house", "a chief's son of Quraysh, and a freed slave from Persia"),
        ("Now", "the banner, in Sālim's ؓ hands")],
        "That was their house. And so, back to the field — where the banner Zayd ؓ dropped is in Sālim's ؓ "
        "hands."),
    "RCT/E-RC23": ("The news reaches Medina", "12 AH", [
        ("al-Yamāma", "the field counted, the dead buried"),
        ("Medina", "the news arrives — ʿUmar ؓ's own brother among the dead"),
        ("Then", "the order that is why you can hold a muṣḥaf")],
        "The field is counted and the dead are buried. Now the news reaches Medina — and ʿUmar ؓ's own "
        "brother is among the dead."),
    "ABU/E-Q7": ("The man he sent for", "12 AH, and 1 AH", [
        ("The dead", "the reciters among them — and no number"),
        ("The order", "gather the Qurʾān"),
        ("Now", "the man he sent for, as a boy")],
        "Abū Bakr ؓ decided the Qurʾān must be gathered, and he sent for one man. Before we hear what he told "
        "him — who was he? Go back with me to the year the Prophet ﷺ came to Medina."),
    "TMW/E-TRN6": ("Meanwhile: the man sent away", "11–12 AH", [
        ("Evening 3", "'ʿIkrima ؓ came into Abyan from Mahra'"),
        ("Last week", "sent on from al-Yamāma: 'do not come back'"),
        ("Tonight", "how he got there — and who he was")],
        "On evening 3 you heard one line: ʿIkrima ؓ came into Abyan from Mahra, with a great many men. Last "
        "week you heard why he was sent on from al-Yamāma: do not come back — go on to Oman and Mahra. "
        "Tonight, how he got to Mahra. But first, who he was."),
    "KTK/E-KD05": ("From Abyan to the last front", "11→12 AH", [
        ("Evening 3", "al-Muhājir ؓ in Ṣanʿāʾ — then 'your own post'"),
        ("Tonight", "ʿIkrima ؓ, from Mahra to Abyan"),
        ("Now", "Ḥaḍramawt and Kinda — the last front")],
        "On evening 3 you left al-Muhājir ؓ in Ṣanʿāʾ, with Abū Bakr's ؓ double order: first the Yemen, then "
        "go on to your own post — Kinda. ʿIkrima ؓ had come into Abyan from Mahra; now you know the road he "
        "came by. From here the two of them go east, to the last front of the war. First: who were Kinda?"),
}

# ---------------------------------------------------------------------------------------------- family trees (#65, #67)
# before card id -> a tree. Boxes: (key, label, centre x, top y, width, style) in inches; style is "focus" (the
# men the part is about), "house" (the one the house is named for), "plain", or "other" (the other side at Badr).
# Links join boxes BY KEY and are drawn as PowerPoint connectors glued to the boxes (#67), so Daniyal can move a
# box and its lines follow: ("child", parent, child) is an elbow from the parent's foot to the child's head;
# ("marriage", left, right) a gold line between two neighbours; ("freed", a, b, caption) a dashed grey line.
# Every link drawn is one the runsheet's tree table cites; the notes carry one sentence per person, each a
# card's or a cited page's.
H_BOX = 0.78
TREES = {
    "THO/E-HS1": dict(
        head="The house of Umm Sulaym ؓ", kick="Banū al-Najjār, of Medina",
        boxes=[("nadr", "al-Naḍr", 2.58, 1.42, 1.70, "plain"), ("milhan", "Milḥān", 9.60, 1.42, 1.70, "plain"),
               ("anas_nadr", "Anas b. al-Naḍr ؓ", 1.55, 2.95, 1.86, "plain"), ("malik", "Mālik", 3.61, 2.95, 1.86, "plain"),
               ("umm_sulaym", "Umm Sulaym ؓ", 5.67, 2.95, 1.86, "house"), ("abu_talha", "Abū Ṭalḥa ؓ", 7.73, 2.95, 1.86, "plain"),
               ("umm_haram", "Umm Ḥarām ؓ", 9.79, 2.95, 1.86, "plain"), ("haram", "Ḥarām ؓ", 11.80, 2.95, 1.70, "plain"),
               ("bara", "al-Barāʾ ؓ", 2.40, 5.15, 1.86, "focus"), ("anas", "Anas ؓ", 4.64, 5.15, 1.86, "plain"),
               ("son", "a small son", 6.70, 5.15, 1.86, "plain")],
        links=[("child", "nadr", "anas_nadr"), ("child", "nadr", "malik"),
               ("child", "milhan", "umm_sulaym"), ("child", "milhan", "umm_haram"), ("child", "milhan", "haram"),
               ("marriage", "malik", "umm_sulaym"), ("marriage", "umm_sulaym", "abu_talha"),
               ("child", "malik", "bara"), ("child", "malik", "anas"), ("child", "umm_sulaym", "anas"),
               ("child", "umm_sulaym", "son"), ("child", "abu_talha", "son")],
        legend=[("marriage", "married", 8.9, 6.35)],
        notes=(
            "TREE — the house of Umm Sulaym ؓ. Point as you go. (Every line is glued to its boxes: move a box and "
            "its lines follow.)\n\n"
            "SAY — Before we go on with him, look at the house he came from: one house of Banū al-Najjār, in "
            "Medina.\n"
            "1. Mālik b. al-Naḍr — the father of al-Barāʾ ؓ and of Anas ؓ. Al-Barāʾ ؓ is the elder.\n"
            "2. Umm Sulaym ؓ — Anas's ؓ mother, and Mālik's wife. When Mālik died, Abū Ṭalḥa ؓ asked for her "
            "hand — the next slide is what she asked of him.\n"
            "3. Anas ؓ — she brought him to the Prophet ﷺ at ten, dressed in half her own head-covering, to "
            "serve him; ten years, and he was never once told off.\n"
            "4. And their small son — the second slide from here.\n\n"
            "IF THE ROOM IS WITH YOU — one line each, pointing:\n"
            "5. Anas b. al-Naḍr ؓ — their uncle, the man Anas ؓ was named after. At Uḥud he said: I find the "
            "scent of Paradise, on this side of the mountain. He was found with more than eighty wounds, and "
            "only his sister knew him — by his fingertips. Nine years later his nephew al-Barāʾ ؓ came off the "
            "wall at al-Yamāma with eighty-odd wounds, and lived.\n"
            "6. Ḥarām ؓ — her brother. Sent to Najd to teach, speared from behind under safe-conduct, he said: "
            "I have won, by the Lord of the Kaʿba. The Prophet ﷺ went into no house in Medina but hers, because "
            "her brother had been killed alongside him.\n"
            "7. Umm Ḥarām ؓ — her sister, who asked to be among the first of this ummah to ride the sea, and was "
            "told she would be; years later, in a caliphate we have not reached, she was.\n"
            "8. Abū Ṭalḥa ؓ — Bayruḥāʾ, the palm-garden he loved most, he gave away. And at Ḥunayn, when the "
            "line was driven back, Umm Sulaym ؓ stood with a dagger: if any of them comes near me, I will use it "
            "on him.\n\n"
            "⚠ Our pages name Umm Sulaym ؓ as Anas's ؓ mother; none names al-Barāʾ's ؓ mother. Say \"Anas's "
            "mother\" and \"his elder brother al-Barāʾ ؓ\" — the tree draws only what the pages say.\n"
            "⚠ Umm Ḥarām ؓ: do not name the fleet's commander. Ḥunayn: stop where the card stops."),
        background=["THO/E-HS2", "THO/E-HS3", "THO/E-HS4", "THO/E-HS5", "THO/E-HS6", "THO/E-HS8", "THO/E-HS11",
                    "THO/E-HS12", "THO/E-HS13", "THO/E-HS10", "THO/E-HS7"]),
    "AHA/E-AS02": dict(
        head="The house of ʿUtba b. Rabīʿa", kick="Quraysh — Badr, 2 AH · al-Yamāma, 12 AH",
        boxes=[("suhayl", "Suhayl b. ʿAmr ؓ", 1.90, 1.42, 2.50, "plain"), ("rabia", "Rabīʿa", 5.50, 1.42, 1.70, "plain"),
               ("utba", "ʿUtba b. Rabīʿa", 4.30, 2.95, 2.40, "other"), ("shayba", "Shayba", 6.70, 2.95, 1.90, "other"),
               ("sahla", "Sahla ؓ", 1.90, 4.45, 1.90, "plain"), ("abu_hudhayfa", "Abū Ḥudhayfa ؓ", 4.30, 4.45, 2.40, "focus"),
               ("walid", "al-Walīd", 6.70, 4.45, 1.90, "other"), ("hind", "Hind ؓ", 8.95, 4.45, 1.80, "plain"),
               ("abu_sufyan", "Abū Sufyān ؓ", 11.25, 4.45, 2.55, "plain"),
               ("salim", "Sālim ؓ", 4.30, 6.00, 1.90, "focus"), ("muawiya", "Muʿāwiya ؓ", 10.075, 6.00, 2.10, "plain")],
        links=[("child", "rabia", "utba"), ("child", "rabia", "shayba"),
               ("child", "utba", "abu_hudhayfa"), ("child", "utba", "walid"), ("child", "utba", "hind"),
               ("child", "suhayl", "sahla"), ("marriage", "sahla", "abu_hudhayfa"),
               ("marriage", "hind", "abu_sufyan"), ("child", "hind", "muawiya"), ("child", "abu_sufyan", "muawiya"),
               ("freed", "abu_hudhayfa", "salim", "freed slave")],
        legend=[("other", "killed at Badr, 2 AH", 8.10, 1.42), ("focus", "killed at al-Yamāma, 12 AH", 8.10, 1.90),
                ("marriage", "married", 8.10, 2.38)],
        notes=(
            "TREE — the house of ʿUtba b. Rabīʿa, and who married into it. Point as you go. (Every line is glued "
            "to its boxes: move a box and its lines follow.)\n\n"
            "SAY — One family, on two battlefields — and some names you know.\n"
            "1. ʿUtba b. Rabīʿa — al-Dhahabī calls him the elder of the Jāhiliyya — his brother Shayba, and his "
            "son al-Walīd: the three who came out for the single combat that opened Badr. All three were killed.\n"
            "2. Abū Ḥudhayfa ؓ, ʿUtba's son, stood with the Muslims that day: his father and his brother on one "
            "side, himself on the other.\n"
            "3. His sister: Hind bint ʿUtba ؓ — al-Dhahabī names her as Umm Muʿāwiya. Her husband, Abū Sufyān ؓ: "
            "the head of Quraysh, their commander at Uḥud and at the Trench — until, in al-Dhahabī's words, "
            "Allah rescued him with Islam on the day of the Conquest. Their son, Muʿāwiya ؓ — a name for "
            "evenings we have not reached.\n"
            "4. His wife: Sahla ؓ, the daughter of Suhayl b. ʿAmr ؓ — the orator of Quraysh, the man who held "
            "Mecca for Islam when the Prophet ﷺ died. She went with Abū Ḥudhayfa ؓ on the first ship to "
            "Abyssinia — hired, it is reported, for half a dīnār.\n"
            "5. Sālim ؓ — the freed slave the household kept. His origin was Iṣṭakhr, in Persia.\n"
            "6. Ten years after Badr, Abū Ḥudhayfa ؓ and Sālim ؓ fell at al-Yamāma, side by side.\n\n"
            "IF THE ROOM IS WITH YOU — one line each:\n"
            "7. Abū Ḥudhayfa ؓ married Sālim ؓ to his own brother al-Walīd's daughter — into the family itself. "
            "Ibn Kathīr alone carries it.\n"
            "8. Another brother, Abū Hāshim b. ʿUtba ؓ — a Muslim on the day of the Conquest, \"and his Islam was "
            "good\"; Muʿāwiya ؓ came to his bedside in his last illness.\n"
            "9. The Prophet ﷺ went out one night with his cloak to listen to Sālim ؓ recite in the mosque; "
            "al-Dhahabī marks the chain good.\n\n"
            "⚠ Abū Sufyān ؓ and Muʿāwiya ؓ are forward references: say \"names for evenings we have not "
            "reached\", and nothing of what came later.\n"
            "⚠ The verses Hind is said to have spoken against her brother at Badr are not to be read out.\n"
            "⚠ Abū Ḥudhayfa's ؓ son Muḥammad, born in Abyssinia, is not on the tree: his later life belongs to "
            "the مشاجرات and is not for the platform.\n"
            "⚠ The tree says \"freed slave\", not \"adopted\": the adoption, and what the law did with it, is not "
            "for the platform."),
        background=["AHA/E-AS03", "AHA/E-AS01", "AHA/E-AS08", "AHA/E-AS05"]),
    "TMW/E-TRN7": dict(
        head="ʿIkrima ؓ and Umm Ḥakīm ؓ", kick="Banū Makhzūm, of Mecca",
        boxes=[("hisham", "Hishām b. al-Mughīra", 6.55, 1.42, 2.90, "plain"),
               ("abu_jahl", "Abū Jahl\nkilled at Badr", 4.20, 2.95, 2.60, "other"),
               ("harith", "al-Ḥārith b. Hishām ؓ", 8.90, 2.95, 2.60, "plain"),
               ("ikrima", "ʿIkrima ؓ", 4.20, 5.15, 2.60, "focus"), ("umm_hakim", "Umm Ḥakīm ؓ", 8.90, 5.15, 2.60, "focus")],
        box_h=0.92,
        links=[("child", "hisham", "abu_jahl"), ("child", "hisham", "harith"),
               ("child", "abu_jahl", "ikrima"), ("child", "harith", "umm_hakim"), ("marriage", "ikrima", "umm_hakim")],
        legend=[("marriage", "married", 8.9, 6.60)],
        notes=(
            "TREE — whose son, and whose daughter. Point as you go. (Every line is glued to its boxes: move a box "
            "and its lines follow.)\n\n"
            "SAY — Whose son he was, and whose daughter his wife was.\n"
            "1. Abū Jahl — who led Quraysh against the Prophet ﷺ at Badr, and was killed there.\n"
            "2. His brother, al-Ḥārith b. Hishām ؓ — in al-Dhahabī's words, one of the noble Companions.\n"
            "3. So Umm Ḥakīm ؓ — al-Ḥārith's daughter, and ʿIkrima's ؓ wife — was his cousin: his father's "
            "brother's daughter. On the day of the Conquest she became a Muslim. He ran.")),
    "KTK/E-KD05": dict(
        head="Kinda", kick="kings once — and two branches in 11 AH",
        boxes=[("akil", "Ākil al-Murār", 2.20, 2.55, 2.90, "plain"),
               ("muawiya_k", "Banū Muʿāwiya", 8.60, 2.10, 2.70, "plain"),
               ("amr", "Banū ʿAmr", 6.90, 3.75, 2.40, "plain"), ("harith_k", "Banū al-Ḥārith", 10.40, 3.75, 2.60, "plain"),
               ("kings", "the four kings", 6.90, 5.20, 2.60, "other"), ("ashath", "al-Ashʿath b. Qays", 10.40, 5.20, 2.90, "focus")],
        links=[("child", "muawiya_k", "amr"), ("child", "muawiya_k", "harith_k"),
               ("child", "amr", "kings"), ("child", "harith_k", "ashath")],
        rules=[("dash", 4.30, 1.85, 4.30, 6.70)],
        captions=[("kings of the north, before Islam", 2.20, 3.55, 3.30)],
        notes=(
            "TREE — Kinda. Point as you go.\n\n"
            "SAY — Who were Kinda?\n"
            "1. A tribe of the south, of Qaḥṭān, that had once given kings to the northern Arabs: the tribes of "
            "Bakr, tired of the strong eating the weak, asked a Tubbaʿ of Yemen for a king, and he gave them a man "
            "of Kinda — Ḥujr b. ʿAmr, Ākil al-Murār.\n"
            "2. Its height: al-Ḥārith b. ʿAmr, given al-Ḥīra by one Persian king and hunted out of it by the next; "
            "forty-eight men of the house were put to death. Ibn al-Athīr says openly that some of this looks like "
            "Kinda praising itself.\n"
            "3. In 11 AH, Kinda's Banū Muʿāwiya had two branches: Banū ʿAmr, whose four brothers the books call "
            "the four kings — and Banū al-Ḥārith, the branch of al-Ashʿath b. Qays.\n"
            "4. Side by side, with no line between them: the books show the royal past and the royal title, and do "
            "not say that one caused the other. Nor do we.\n\n"
            "IF THE ROOM IS WITH YOU: one of al-Ḥārith's sons, another Ḥujr, ruled Banū Asad harshly — and they "
            "killed him in his tent.\n\n"
            "⚠ No date before Islam is on this slide, and none is to be said: they are estimates."),
        background=["KTK/E-KD01", "KTK/E-KD02", "KTK/E-KD03", "KTK/E-KD06"]),
}

# ---------------------------------------------------------------------------------------------- checkpoints
# after card id -> (Line png, map (scene, step), keys, the SAY lines, the time check, CARRYING ON,
#                   next week's question, three lesson cards)
CHECKPOINTS = {
    "RCT/E-RC37": dict(
        n=1, line="line_s05_cp1.png", scene=("s05-09-arabia", 1),
        keys=[("al-Yamāma", "taken — and its dead"), ("The east", "still fighting")],
        say_line="Tonight so far: the terms at the forts; the two houses and the men beside them; the reciters "
                 "— and the order at Medina that they are the reason for.",
        say_map="When you came in, al-Yamāma had just been taken. Now you know what it cost. And look at the "
                "east: still fighting.",
        clock="Past 0:40 here → close here. ʿIkrima's ؓ road — Oman, Mahra and the last front — opens evening 6, "
              "whole.",
        carry="While all of this was happening, one commander had been sent away from al-Yamāma, and told not "
              "to come back.",
        question="One commander was sent away from al-Yamāma before the battle and told not to come back. Next "
                 "week: where his road went — all the way to the last front of the war.",
        lessons=["RCT/E-RC67", "THO/E-HS9", "RCT/E-RC37"]),
    "RCT/E-RC30": dict(
        n=2, line="line_s05_cp2.png", scene=("s05-09-arabia", 2),
        keys=[("Oman, Mahra", "settled"), ("Ḥaḍramawt", "still to come")],
        say_line="Tonight: what al-Yamāma cost, the Qurʾān gathered — and ʿIkrima's ؓ road, through Oman and "
                 "Mahra.",
        say_map="When you came in, the east was burning. Oman and Mahra are settled. One front is left: "
                "Ḥaḍramawt.",
        clock="Past 0:40 here → close here. Ḥaḍramawt opens evening 6.",
        carry="One front was left — and it began with a quarrel about camels.",
        question="One front was left, and the books say it began with one she-camel. Next week: how.",
        lessons=["RCT/E-RC37", "RCT/E-RC29", "RCT/E-RC30"]),
    "TSY/E-YK12": dict(
        n=3, line="line_s05_cp3.png", scene=("s05-09-arabia", 3),
        keys=[("al-Nujayr", "opened"), ("al-Ashʿath", "in bonds, for Medina")],
        say_line="Tonight: al-Yamāma's cost, the Qurʾān, Oman and Mahra — and now Kinda, from a camel to a fort.",
        say_map="When you came in, the east was burning. Now every front we have told is settled — and one man "
                "is on his way to Medina in bonds.",
        clock="Past 0:42 here → close here. The judgement at Medina opens evening 6.",
        carry="He had left his own name off the paper. Now he stood in front of Abū Bakr ؓ.",
        question="A man who had led a whole tribe out of Islam stood in front of Abū Bakr ؓ, in bonds. Next "
                 "week: what Abū Bakr ؓ did with him.",
        lessons=["RCT/E-RC37", "RCT/E-RC30", "TSY/E-YK12"]),
}

CLOSE_AFTER = "POT/E-PG41"
TONIGHT = [("How the day ended", "RCT/E-RC67"), ("The house of Umm Sulaym ؓ", "THO/E-HS9"),
           ("The Qurʾān gathered", "RCT/E-RC37"), ("ʿIkrima's ؓ road", "RCT/E-RC30"),
           ("The last front", "POT/E-PG41")]
NEXT_WEEK = ("Sixteen riders left Medina for Bahrayn. By the time they reached it they were an army. How?")

# ---------------------------------------------------------------------------------------------- notes
REF = re.compile(r"\(\s*`[^`]*`[^)]*\)|`[^`]*\.md`[^,;)]*|\(#\d+[^)]*\)|\s*#\d+\b|\(§[\d.]+[^)]*\)|§[\d.]+|\(\s*\)")
SKIP = re.compile(r"SKIP BEATS\s*([\d,\s]+)")


HOUSEKEEPING = re.compile(r"\*{0,2}New card\*{0,2}[.,]?\s*|Tier raised from GOOD:[^.|]*\.?\s*|"
                          r"from the campaign note\s*[§\d.]*:?\s*|\(\s*\)")


def runsheet_note(note):
    """The runsheet's remark on the card, fit for the pane: no file names, section or decision numbers, and
    none of the runsheet's own housekeeping ("New card", "Tier raised") — that is for the runsheet reader."""
    # an inline Arabic citation («الکامل ج۲ ص۲۱۸») comes out reversed in the notes pane once its digits are
    # Latin; say it in English, as the slide faces do (build_full_deck.cite_en)
    def iso(m):
        s = m.group(1)
        if re.search(r"[جص]\s*[0-9۰-۹٠-٩]", s):
            return B.cite_en(s)
        r = B.romanise(s)                        # a book's name alone goes to English, as the faces do
        return r if not B.ARABIC_CH.search(r) else s
    t = re.sub("⁨([^⁩]*)⁩", iso, note or "")
    # "SKIP BEATS 2 — the card before told them": the instruction is the build's, the reason the runsheet
    # reader's; neither is a line for the lectern
    t = re.sub(r"SKIP BEATS\s*[\d,\s]+(?:\s*—\s*[^.|]*\.?)?", "", t)
    t = HOUSEKEEPING.sub("", REF.sub("", t))
    t = B.clean(t)
    t = re.sub(r"\s*\.\s*\.", ".", t)
    t = re.sub(r"\s{2,}", " ", t).strip(" .·—")
    return ("NOTE — " + t) if t else ""


def card_notes(c, note):
    """#60: warnings first, SAY, the quotation, the lesson line, hands-up; background last."""
    skip = set()
    m = SKIP.search(note or "")
    if m:
        skip = {int(x) for x in re.findall(r"\d+", m.group(1))}
    c = dict(c)
    if skip and c["beats"]:
        kept = [(i, b) for i, b in enumerate(c["beats"], 1) if i not in skip]
        qa = c.get("quote_after")
        c["beats"] = [b for _, b in kept]
        if qa:
            c["quote_after"] = sum(1 for i, _ in kept if i <= qa) or None
    c["map"] = None                              # the map's description is build apparatus, not a line to say
    body = B.notes_for(c)
    # the ⚠ blocks notes_for carries go to the top
    warn = [b for b in body.split("\n\n") if b.startswith("⚠")]
    rest = [b for b in body.split("\n\n") if not b.startswith("⚠")]
    top = [x for x in [runsheet_note(note)] + warn if x]
    # markdown bold/code markers and the editor's bidi isolates are noise in PowerPoint's notes pane
    return re.sub("\\*\\*|`|[⁨⁩]", "", "\n\n".join(top + rest))


def map_card_notes(c, note, clicks, words):
    """#63: the card told on its own map. Its beats carry ▶ CLICK n on the beat the map moves on; when the card
    has a quotation for the face, the quotation and the lesson line go to the words slide after the map, and
    SAY ends by sending the speaker there. A hands-up that belongs before the telling goes to the top."""
    skip = set()
    m = SKIP.search(note or "")
    if m:
        skip = {int(x) for x in re.findall(r"\d+", m.group(1))}
    kept = [i for i in range(1, len(c["beats"]) + 1) if i not in skip]
    for b, _ in clicks:
        if b not in kept:
            raise SystemExit("%s: a click sits on beat %d, which the runsheet skips or the card lacks" % (c["id"], b))
    cm = dict(c)
    if words:
        cm["arabic"] = None
        cm["ibrah"] = ""
    blocks = card_notes(cm, note).split("\n\n")
    out, n = [], 0
    hands = [b for b in blocks if b.startswith("HANDS UP")]
    for b in blocks:
        if b in hands:
            continue
        if b.startswith("SAY —"):
            lines = []
            for ln in b.split("\n"):
                mm = re.match(r"^(\d+)\. ", ln)
                if mm:
                    orig = kept[int(mm.group(1)) - 1]
                    for bb, txt in clicks:
                        if bb == orig:
                            n += 1
                            lines.append("▶ CLICK %d — %s" % (n, txt))
                lines.append(ln)
            if words:
                lines.append("▶ NEXT SLIDE — the words: read the Arabic, then its English.")
            b = "\n".join(lines)
            out.extend("HANDS UP FIRST — ask it before the first click. "
                       + re.sub(r"^(?:Yes|yes)\b[\s.,;:—–-]*", "", h[len("HANDS UP — "):]) for h in hands)
        out.append(b)
    return "\n\n".join(out), n


def words_notes(c):
    """The slide after a card's map: the quotation and its lesson line, nothing told twice (#63)."""
    body = B.notes_for(dict(c, beats=[], what="", hands="no", extra="", map=None))
    return re.sub("\\*\\*|`|[⁨⁩]", "", "THE WORDS — the story was told on the map. Read the Arabic, then its "
                                       "English — and the one line.\n\n" + body)


def map_notes(intro, clicks, hands=None):
    lines = []
    if hands:
        lines.append("HANDS UP FIRST — ask it before the first click:\n" + hands)
    lines.append("SAY — " + intro)
    for i, l in enumerate(clicks, 1):
        lines.append("▶ CLICK %d — %s" % (i, l))
    return "\n\n".join(lines)


def pool_background(pool, ids):
    """The beats of the cards the second pass moved off the running order, for the tree's BACKGROUND."""
    out = []
    for cid in ids:
        c = pool[cid]
        beats = "; ".join("%s — %s" % (h, t) for h, t in c["beats"]) or B.short(c["what"], 60)
        out.append("• %s: %s" % (B.clean(c["title"]), B.clean(beats)))
    return ("BACKGROUND (read at home, not aloud) — the house's stories in full:\n" + "\n".join(out)) if out else ""


# ---------------------------------------------------------------------------------------------- slide makers
def jpeg(png):
    """A map ground as a quality-90 JPEG beside its PNG. Shaded relief and grain do not compress as PNG —
    the first evening-5 build was 80 MB with PNG grounds. The layers stay PNG: they need their alpha."""
    jpg = os.path.splitext(png)[0] + ".jpg"
    if not os.path.exists(jpg) or os.path.getmtime(jpg) < os.path.getmtime(png):
        from PIL import Image
        Image.open(png).convert("RGB").save(jpg, "JPEG", quality=90, optimize=True)
    return jpg


LAYERED = {}      # slide id -> (scene, first step, last step), for notes_book.py


def layered_map(prs, scene, a, b, headline, kicker, keys):
    stem = os.path.join(MAPS, "%s-r%d-%d" % (scene, a, b))
    base, man = stem + "-base.png", stem + "-layers.json"
    if not (os.path.exists(base) and os.path.exists(man)):
        raise SystemExit("map layers missing for %s %d-%d — run make_maps.py first" % (scene, a, b))
    s = D.map_slide(prs, jpeg(base), headline, keys=keys, kicker=kicker, map_frac=0.70, trim=False)
    placed = anim.place_layers(s, man)
    n = anim.animate(s, placed, prs.slide_width, prs.slide_height)
    LAYERED[s.slide_id] = (scene, a, b)
    return s, n


def static_map(prs, scene, step, headline, kicker, keys):
    png = os.path.join(MAPS, "%s-step-%02d.png" % (scene, step))
    if not os.path.exists(png):
        raise SystemExit("map render missing: %s — run make_maps.py first" % png)
    return D.map_slide(prs, jpeg(png), headline, keys=keys, kicker=kicker, map_frac=0.70, trim=False)


BOX_STYLE = {   # fill, border, border width, text colour, dashed border
    "focus": (D.TEAL, D.TEAL, 1.5, D.CREAM, False),
    "house": (D.CREAM, D.GOLD, 3.0, D.DARK, False),
    "plain": (D.CREAM, D.RULE, 1.5, D.DARK, False),
    "other": (D.WHITE, D.MUTED, 1.5, D.MUTED, True),
}
TOP, LEFT, BOTTOM, RIGHT = 0, 1, 2, 3        # a rectangle's connection sites, in PowerPoint's order


def _link(s, a, b, kind, a_site, b_site):
    """A connector glued to two boxes (#67). PowerPoint re-routes it when either box is moved."""
    ax, ay = {TOP: (a.left + a.width // 2, a.top), BOTTOM: (a.left + a.width // 2, a.top + a.height),
              LEFT: (a.left, a.top + a.height // 2), RIGHT: (a.left + a.width, a.top + a.height // 2)}[a_site]
    bx, by = {TOP: (b.left + b.width // 2, b.top), BOTTOM: (b.left + b.width // 2, b.top + b.height),
              LEFT: (b.left, b.top + b.height // 2), RIGHT: (b.left + b.width, b.top + b.height // 2)}[b_site]
    c = s.shapes.add_connector(MSO_CONNECTOR.ELBOW if kind == "child" else MSO_CONNECTOR.STRAIGHT, ax, ay, bx, by)
    c.begin_connect(a, a_site)
    c.end_connect(b, b_site)
    if kind == "marriage":
        c.line.color.rgb, c.line.width = D.GOLD, Pt(3.0)
    elif kind == "freed":
        c.line.color.rgb, c.line.width = D.MUTED, Pt(2.25)
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        c.line.color.rgb, c.line.width = D.TEAL, Pt(2.25)
    c.name = "%s: %s - %s" % (kind, a.name.split(": ", 1)[-1], b.name.split(": ", 1)[-1])
    return c


def _rule(s, x1, y1, x2, y2):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb, c.line.width = D.MUTED, Pt(2.25)
    c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    c.name = "rule"
    return c


def family_tree(prs, pool, t):
    """#65/#67: a house as an editable org chart. Names and lines on the face; one sentence per person in the
    notes. Boxes are named shapes ("person: Hind"); every line is a connector glued to its two boxes."""
    s = D.blank(prs)
    D.header(s, t["head"], t.get("kick"))
    bh = t.get("box_h", H_BOX)
    box = {}
    for key, label, cx, y, w, style in t["boxes"]:
        fill, border, bw, ink, dashed = BOX_STYLE[style]
        r = D.rect(s, Inches(cx - w / 2), Inches(y), Inches(w), Inches(bh), fill=fill, line_color=border,
                   line_w=Pt(bw), dash=dashed, name="person: " + label.split("\n")[0])
        tf = r.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.06)
        tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(label.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            p.line_spacing = 1.0
            run = p.add_run()
            run.text = line
            run.font.size, run.font.name, run.font.bold = Pt(24), D.EN, True
            run.font.color.rgb = ink
            D.set_cs(run, D.AR_HON if D._HONORIFIC.search(line) else D.EN)
        box[key] = r
    for link in t.get("links", []):
        kind, a, b = link[0], box[link[1]], box[link[2]]
        if kind == "child":
            _link(s, a, b, kind, BOTTOM, TOP)
        elif kind == "marriage":
            _link(s, a, b, kind, RIGHT, LEFT)
        elif kind == "freed":
            c = _link(s, a, b, kind, BOTTOM, TOP)
            if len(link) > 3:
                D.text(s, link[3], Inches((a.left + a.width / 2) / 914400 + 0.12), Inches((a.top + a.height) / 914400 + 0.14),
                       Inches(1.75), Inches(0.42), size=24, color=D.MUTED, font=D.SANS)
    for x1, y1, x2, y2 in [r[1:] for r in t.get("rules", [])]:
        _rule(s, x1, y1, x2, y2)
    for text, cx, y, w in t.get("captions", []):
        D.text(s, text, Inches(cx - w / 2), Inches(y), Inches(w), Inches(0.9), size=24, color=D.MUTED,
               font=D.SANS, align=PP_ALIGN.CENTER, line=1.05)
    for style, text, x, y in t.get("legend", []):
        if style == "marriage":
            c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y + 0.23), Inches(x + 0.42), Inches(y + 0.23))
            c.line.color.rgb, c.line.width = D.GOLD, Pt(3.0)
            c.name = "legend"
        else:
            fill, border, bw, _, dashed = BOX_STYLE[style]
            D.rect(s, Inches(x), Inches(y + 0.06), Inches(0.42), Inches(0.34), fill=fill, line_color=border,
                   line_w=Pt(bw), dash=dashed, name="legend")
        D.text(s, text, Inches(x + 0.56), Inches(y), Inches(4.4), Inches(0.5), size=24, color=D.INK, font=D.SANS)
    bg = pool_background(pool, t.get("background", []))
    D.note(s, t["notes"] + ("\n\n" + bg if bg else ""))
    return s


def lessons_of(pool, ids):
    return [B.clean(pool[i]["ibrah"]) for i in ids]


def checkpoint(prs, pool, cp):
    s1 = D.timeline_slide(prs, os.path.join(VIS, cp["line"]), "Where we stand", kicker="So far tonight")
    D.note(s1, "CHECKPOINT %d — the Line.\n\nSAY — %s\n\n(Next slide: the map. The time check is there.)"
           % (cp["n"], cp["say_line"]))
    s2 = static_map(prs, cp["scene"][0], cp["scene"][1], "Where we stand", "So far tonight", cp["keys"])
    close = "\n".join("   %d. %s" % (i, l) for i, l in enumerate(lessons_of(pool, cp["lessons"]), 1))
    D.note(s2, "CHECKPOINT %d — the map.\n\nSAY — %s\n\n⏱ %s\n\nCARRYING ON → say: \"%s\" — and click on.\n\n"
               "OUT OF TIME → this is your ending. Stay on this map and say:\n%s\n\nThen next week: \"%s\"\n\n"
               "Then a loud السلام علیکم, and the dua. Next week opens on these two slides."
           % (cp["n"], cp["say_map"], cp["clock"], cp["carry"], close, cp["question"]))
    return s1, s2


def the_close(prs, pool):
    s = D.timeline_slide(prs, os.path.join(VIS, "line_s05_close.png"), "Tonight on the Line")
    D.note(s, "THE CLOSE — the Line.\n\nSAY — Tonight began at the forts of al-Yamāma, with what was signed. We "
              "went into two houses — Umm Sulaym's ؓ, and Abū Ḥudhayfa's ؓ and Sālim's ؓ — and came out on the "
              "field where their men fell, and at Medina, where their deaths became the reason the Qurʾān was "
              "gathered. And then one man's road: ʿIkrima ؓ, through Oman and Mahra to the last front.")
    lines = ["Oman and Mahra, settled.",
             "Ḥaḍramawt — the last front — settled.",
             "And the road one man rode through all of it: ʿIkrima ؓ, from al-Yamāma to al-Nujayr."]
    m, n = layered_map(prs, "s05-09-arabia", 1, 4, "Where we stand", "12 AH",
                       [("Tonight", "Oman, Mahra, and the last front"), ("Next week", "Bahrayn, told as 'meanwhile'")])
    D.note(m, map_notes("When you came in, al-Yamāma had just been taken, and the east was burning.", lines))
    if n != len(lines):
        raise SystemExit("the close map has %d clicks; its notes have %d" % (n, len(lines)))
    s = D.lessons_slide(prs, [h for h, _ in TONIGHT])
    D.note(s, "TONIGHT — say each heading, then its line:\n" + "\n".join(
        "%d. %s — %s" % (i, h, B.clean(pool[c]["ibrah"])) for i, (h, c) in enumerate(TONIGHT, 1)))
    q = D.question_slide(prs, NEXT_WEEK)
    D.note(q, "Ask it, and pause. Then a loud السلام علیکم — and the dua. The room must know it has ended.")


# ---------------------------------------------------------------------------------------------- build
def bookend_images():
    """Daniyal's own STOP D pair from evening 4 — his images, not the build's (#23, DELIVERED.md)."""
    os.makedirs(BOOKEND, exist_ok=True)
    line, mp = os.path.join(BOOKEND, "stop_d_line.png"), os.path.join(BOOKEND, "stop_d_map.png")
    if not (os.path.exists(line) and os.path.exists(mp)):
        from pptx import Presentation
        share = Presentation(os.path.join(ROOT, "S04_kinda_butah_yamama", "S04_share.pptx"))
        slides = list(share.slides)
        pics = {}
        for idx, key in ((63, "line"), (64, "map")):              # his slides 67-68 = share 64-65
            p = [sh for sh in slides[idx].shapes if sh.shape_type == 13]
            pics[key] = max(p, key=lambda sh: sh.width * sh.height).image.blob
        open(line, "wb").write(pics["line"])
        open(mp, "wb").write(pics["map"])
    return line, mp


HASH_FILE = os.path.join(HERE, ".build", "deck.sha256")      # what the last build wrote, so it may overwrite its own work


def _sha256(path):
    import hashlib
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def deck_has_uncommitted_changes(path):
    """True when the deck on disk is neither its committed version nor this build's own last output — i.e.
    somebody has saved edits to it since."""
    import subprocess
    if not os.path.exists(path):
        return False
    try:
        r = subprocess.run(["git", "status", "--porcelain", "--", path], capture_output=True, text=True, cwd=ROOT)
    except OSError:
        return False
    if not any(line[:2].strip() in ("M", "AM", "MM") for line in r.stdout.splitlines()):
        return False                                     # identical to the commit: nobody's edits are at risk
    try:
        last = open(HASH_FILE, encoding="utf-8").read().strip()
    except OSError:
        last = ""
    return _sha256(path) != last                         # changed by someone other than the last build


def remember_deck(path):
    os.makedirs(os.path.dirname(HASH_FILE), exist_ok=True)
    with open(HASH_FILE, "w", encoding="utf-8") as f:
        f.write(_sha256(path))


def face_card(c, cid):
    c = dict(c)
    if cid in FACE_TITLE:
        c["title"] = FACE_TITLE[cid]
    if cid in FACE_WHEN:
        c["when"] = FACE_WHEN[cid]
    if cid in NO_HANDS:
        c["hands"] = "no"
    if cid in FACE_CITE:
        c["cite"] = FACE_CITE[cid]
    return c


def build():
    pool = B.pool_cards()
    parts = B.runsheet(os.path.join(HERE, "RUNSHEET.md"))
    ids = [cid for _, rows in parts for cid, _ in rows]
    missing = [cid for cid in ids if cid not in pool]
    if missing:
        raise SystemExit("RUNSHEET.md names cards that are not in any pool: %s" % missing)
    for table, name in ((MAP_FOR, "MAP_FOR"), (TREES, "TREES"), (BRIDGE_BEFORE, "BRIDGE_BEFORE"),
                        (CHECKPOINTS, "CHECKPOINTS"), (RECAP_BEFORE, "RECAP_BEFORE")):
        stray = [k for k in table if k not in ids]
        if stray:
            raise SystemExit("%s names cards the runsheet does not run: %s" % (name, stray))
    check_face_quotes(pool)

    prs = D.deck()
    D.title_slide(prs, "What al-Yamāma cost", "12 AH", "Evening 5")

    line, mp = bookend_images()
    s = D.timeline_slide(prs, line, "Where we stopped")
    D.note(s, "BOOKEND IN — last week's closing Line, unchanged.\n\nSAY — Last week ended with al-Yamāma taken: "
              "the garden, the gate, Musaylima dead, and the treaty kept against the caliph's letter.")
    s = D.map_slide(prs, mp, "Where we stand", keys=[("al-Yamāma", "taken"), ("Ḥaḍramawt", "the last front")],
                    kicker="From last week", map_frac=0.70)
    D.note(s, "BOOKEND IN — last week's closing map, unchanged.\n\nSAY — Tonight: what was signed at the forts, "
              "what that day cost — and where the man sent away from al-Yamāma before the battle had gone.")

    made = clicks_total = bridges = maps = words_slides = trees = checkpoints = 0
    for part, rows in parts:
        D.section_slide(prs, B.clean(part.split(" — ")[0], translit=True),
                        B.clean(part.split(" — ", 1)[1] if " — " in part else "", translit=True) or None)
        for cid, rnote in rows:
            c = face_card(pool[cid], cid)
            if cid in RECAP_BEFORE:
                scene, a, b, head, kick, keys, intro, lines = RECAP_BEFORE[cid]
                s, n = layered_map(prs, scene, a, b, head, kick, keys)
                if n != len(lines):
                    raise SystemExit("%s: %d clicks on the slide, %d ▶ CLICK lines in its notes" % (scene, n, len(lines)))
                D.note(s, map_notes(intro, lines))
                maps += 1
                clicks_total += n
            if cid in BRIDGE_BEFORE:
                head, kick, rows_, say = BRIDGE_BEFORE[cid]
                s = D.diagram_slide(prs, head, rows_, kicker=kick)
                D.note(s, "BRIDGE — say:\n" + say)
                bridges += 1
            if cid in TREES:
                family_tree(prs, pool, TREES[cid])
                trees += 1
            off_face = "NOT FOR THE SLIDE FACE" in (c.get("extra") or "")
            if cid in MAP_FOR:
                scene, a, b, keys, clicks = MAP_FOR[cid]
                fq = FACE_QUOTE.get(cid)
                words = bool(c["arabic"]) and not off_face and B.ar_len(fq[0] if fq else c["arabic"]) <= B.AR_SLIDE_MAX
                head = B.headline_for(B.short(B.face(c["title"]), 9))
                s, n = layered_map(prs, scene, a, b, head, B.kicker_for(c), keys)
                notes, k = map_card_notes(c, rnote, clicks, words)
                if n != k:
                    raise SystemExit("%s on %s %d-%d: %d clicks on the slide, %d ▶ CLICK lines in its notes"
                                     % (cid, scene, a, b, n, k))
                D.note(s, notes)
                maps += 1
                clicks_total += n
                if words:
                    w = B.card_slide(prs, c, face_quote=fq)
                    if w is not None:
                        D.note(w, words_notes(c))
                        words_slides += 1
                made += 1
            else:
                face_text = B.short(c["what"].split(". ")[0].rstrip(".") + ".", 18) if off_face else None
                s = B.card_slide(prs, c, arabic_on_face=not off_face, face_text=face_text,
                                 face_quote=FACE_QUOTE.get(cid))
                if s is not None:
                    D.note(s, card_notes(c, rnote))
                    made += 1
            if cid in CHECKPOINTS:
                checkpoint(prs, pool, CHECKPOINTS[cid])
                checkpoints += 1
            if cid == CLOSE_AFTER:
                the_close(prs, pool)

    if checkpoints != len(CHECKPOINTS):
        raise SystemExit("%d checkpoints emitted of %d — a checkpoint card left the runsheet"
                         % (checkpoints, len(CHECKPOINTS)))
    target = os.path.join(HERE, "S05.pptx")
    if deck_has_uncommitted_changes(target):
        # 2026-10-01: a build overwrote S05.pptx while it carried Daniyal's own saved edits, which are not
        # recoverable. A deck that differs from its committed version is his, until he commits it.
        print("   !! S05.pptx differs from its committed version (Daniyal's edits?) - NOT overwritten.\n"
              "      This build is written to S05_NEW.pptx. Commit or discard the changes to S05.pptx first.")
        target = os.path.join(HERE, "S05_NEW.pptx")
    out = D.save(prs, target)
    remember_deck(out)
    import json
    os.makedirs(os.path.join(HERE, ".build"), exist_ok=True)
    with open(os.path.join(HERE, ".build", "slides.json"), "w", encoding="utf-8") as f:
        json.dump({str(i): LAYERED[s.slide_id] for i, s in enumerate(prs.slides, 1) if s.slide_id in LAYERED},
                  f, ensure_ascii=False, indent=1)
    print("   %d of %d runsheet cards · %d maps (%d clicks) · %d words slides · %d trees · %d bridges · "
          "%d checkpoints · %d slides in all"
          % (made, len(ids), maps, clicks_total, words_slides, trees, bridges, checkpoints, len(prs.slides)))
    return out


if __name__ == "__main__":
    if os.path.exists(os.path.join(HERE, "S05.FINAL")):
        raise SystemExit("S05.FINAL exists: S05.pptx is hand-finished. Do not rebuild over it.")
    build()
